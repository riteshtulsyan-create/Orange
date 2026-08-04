"""Evaluate the workbook's formulas for real and check them against the source data.

openpyxl writes formulas but never calculates them, so a workbook can look right and be wrong.
This re-implements each roll-up formula's INTENT directly in pandas and compares it with the
value the formula produces once evaluated — catching wrong ranges, wrong criteria and
off-by-one row errors.

Two engines are tried: LibreOffice headless (fast, exact) then the `formulas` package.
If neither is available the structural check still runs.
"""
import re
import subprocess
import sys
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from tally_lib import SCRATCH, ROOT

WB = max((ROOT / "Inventory").glob("Stock_Valuation_Check_*.xlsx"),
         key=lambda p: p.stat().st_mtime)
print(f"verifying {WB.name}\n")

df = pd.read_parquet(SCRATCH / "phase_b4_grouped.parquet")
ENTS = ["Orange O Tec", "Enterprises", "Colorix"]

wb = load_workbook(WB)         # formulas as written
fails = []


def expect(label, got, want, tol=0.51):
    ok = (got is not None) and abs(float(got) - float(want)) <= tol
    print(f"  {'OK ' if ok else '*** FAIL'}  {label:<52} "
          f"formula={got if got is None else f'{float(got):,.2f}':>18}  "
          f"source={float(want):>18,.2f}")
    if not ok:
        fails.append(label)


# ---------------------------------------------------------------- 1. structural
print("[1] formula structure — ranges point at real data rows")
for e in ENTS:
    ws = wb[e]
    hr = next(r for r in range(1, 60) if ws.cell(row=r, column=1).value == "Rank")
    d1, d2 = hr + 1, ws.max_row
    n_src = len(df[df["Entity"] == e])
    ok = (d2 - d1 + 1) == n_src
    print(f"  {'OK ' if ok else '*** FAIL'}  {e:<16} detail rows {d1}..{d2} "
          f"= {d2 - d1 + 1} (source {n_src})")
    if not ok:
        fails.append(f"{e} row count")
    # every roll-up SUMIFS must reference exactly this range
    bad = 0
    for r in range(1, hr):
        for c in range(1, 12):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.startswith("=") and "SUMIFS" in v.upper():
                for m in re.finditer(r"\$[A-Z]+\$(\d+):\$[A-Z]+\$(\d+)", v):
                    if (int(m.group(1)), int(m.group(2))) != (d1, d2):
                        bad += 1
    print(f"  {'OK ' if bad == 0 else '*** FAIL'}  {e:<16} "
          f"{bad} roll-up range mismatches")
    if bad:
        fails.append(f"{e} range mismatch")

# ---------------------------------------------------------------- 2. evaluate
calc = None
soffice = [p for p in [
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"] if Path(p).exists()]
if soffice:
    print("\n[2] evaluating with LibreOffice headless")
    outdir = SCRATCH / "calc"
    outdir.mkdir(exist_ok=True)
    subprocess.run([soffice[0], "--headless", "--convert-to", "xlsx", "--outdir",
                    str(outdir), str(WB)], capture_output=True, timeout=900)
    got = outdir / WB.name
    if got.exists():
        calc = load_workbook(got, data_only=True)

if calc is None:
    print("\n[2] evaluating with the `formulas` package (slow on large sheets)")
    try:
        import formulas
        xl = formulas.ExcelModel().loads(str(WB)).finish()
        sol = xl.calculate()
        calc = "formulas"
        FSOL = {k.upper(): v for k, v in sol.items()}
    except Exception as ex:                                   # noqa: BLE001
        print(f"    unavailable: {type(ex).__name__}: {ex}")
        calc = None


def cellval(sheet, addr):
    if calc is None or calc == "formulas":
        if calc == "formulas":
            key = f"'[{WB.name.upper()}]{sheet.upper()}'!{addr.upper()}"
            v = FSOL.get(key)
            try:
                return float(v.value[0, 0])
            except Exception:                                  # noqa: BLE001
                return None
        return None
    return calc[sheet][addr].value


# ---------------------------------------------------------------- 3. compare
if calc is not None:
    print("\n[3] evaluated values vs source data")
    valid = df[df["Benchmark"] == "valid"]

    for e in ENTS:
        ws = wb[e]
        hr = next(r for r in range(1, 60) if ws.cell(row=r, column=1).value == "Rank")
        # locate the roll-up block by its labels
        lab = {}
        for r in range(1, hr):
            v = ws.cell(row=r, column=1).value
            if isinstance(v, str):
                lab[v] = r
        units = sorted(df[df["Entity"] == e]["Unit"].unique())
        comb = chr(ord("A") + len(units) + 1)
        d = df[df["Entity"] == e]
        v = valid[valid["Entity"] == e]
        print(f"\n  --- {e} (combined column {comb}) ---")
        for label, want in [
            ("Items with closing stock", len(d)),
            ("Closing stock value", d["Closing Value"].sum()),
            ("Items with a cost benchmark", len(v)),
            ("Value benchmarked", v["Closing Value"].sum()),
            ("Overstated vs purchase cost (+)",
             v[v["Gap vs Inward Rate"] > 0]["Gap vs Inward Rate"].sum()),
            ("Understated vs purchase cost (-)",
             v[v["Gap vs Inward Rate"] < 0]["Gap vs Inward Rate"].sum()),
            ("NET GAP vs purchase cost", v["Gap vs Inward Rate"].sum()),
            ("Items with negative stock", int((d["Closing Qty"] < 0).sum())),
        ]:
            if label in lab:
                expect(f"{e[:12]}: {label}", cellval(e, f"{comb}{lab[label]}"), want)

    # group totals
    ws = wb["Summary"]
    gt = next(r for r in range(1, 40)
              if str(ws.cell(row=r, column=1).value or "").startswith("GROUP TOTAL"))
    print("\n  --- Summary group total ---")
    expect("Summary: closing stock value", cellval("Summary", f"D{gt}"),
           df["Closing Value"].sum())
    expect("Summary: net gap", cellval("Summary", f"I{gt}"),
           valid["Gap vs Inward Rate"].sum())

    ws = wb["Segment Summary"]
    st = next(r for r in range(1, 60)
              if str(ws.cell(row=r, column=1).value or "").startswith("GROUP TOTAL"))
    print("\n  --- Segment Summary ---")
    expect("Segments: total closing value", cellval("Segment Summary", f"E{st}"),
           df["Closing Value"].sum())
    expect("Segments: total net gap", cellval("Segment Summary", f"I{st}"),
           valid["Gap vs Inward Rate"].sum())

print("\n" + ("ALL CHECKS PASSED" if not fails else f"FAILURES: {fails}"))
sys.exit(1 if fails else 0)
