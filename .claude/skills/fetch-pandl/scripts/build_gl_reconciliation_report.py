"""
Compares two dated PL_BS_GL_Detail_<timestamp>.xlsx files (the previous refresh's output vs
this refresh's freshly rebuilt one) and produces a dated
Files/Profit_Change_Reconciliation_<timestamp>.xlsx (timestamp matches the new file):

  1. "PAT Reconciliation"  - Turnover / Gross Profit / PAT bridged old -> new, per company
     and Grand Total, with the specific GL ledger(s) driving each company's PAT swing
     auto-identified from the GL-level diff (not hand-curated).
  2. "GL Changes - PL"     - every P&L ledger Added / Removed / Changed / Renamed between
  3. "GL Changes - BS"       the two files, for each of the 5 companies individually.

"Renamed" detection collapses Tally's habit of returning the same ledger with slightly
different punctuation between pulls (e.g. "(SPARE)" vs "-SPARE") into one row instead of
a fake add + fake remove pair - see SKILL.md "Standing procedure" for why this matters.

Run standalone (defaults auto-discover the two most recent dated files in Files/):
    python build_gl_reconciliation_report.py [--old PATH] [--new PATH] [--out PATH]

This is Step E of the standing CFO-mode refresh procedure (see SKILL.md "File naming and
version history"). It runs every time the workbook is rebuilt, BEFORE the old file gets moved
into the "Old Version of PL BS" archive folder (Step F) - so by default it just compares the
two newest PL_BS_GL_Detail_*.xlsx files still sitting in Files/.
"""
import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BASE_NAME = "PL_BS_GL_Detail"
RECON_BASE_NAME = "Profit_Change_Reconciliation"
TIMESTAMP_FORMAT = "%Y-%m-%d_%H%M"


def find_pl_bs_files(files_dir):
    return sorted(files_dir.glob(f"{BASE_NAME}_*.xlsx"))


def extract_timestamp(path):
    """Pulls the '<timestamp>' out of 'PL_BS_GL_Detail_<timestamp>.xlsx'."""
    stem = Path(path).stem
    return stem[len(BASE_NAME) + 1:] if stem.startswith(BASE_NAME + "_") else None

FILES_DIR = Path(__file__).resolve().parents[4] / "Files"

COMPANY_DEFS = [
    ("Orange O Tec Pvt Ltd", "ORANGE O TEC PRIVATE LIMITED (01-04-25TO31-03-27)"),
    ("Orange O Tec NOIDA", "ORANGE O TEC PRIVATE LIMITED-NOIDA-(from 1-Apr-25)"),
    ("Enterprises Pvt Ltd", "ORANGE O TEC ENTERPRISES PVT LTD(F.Y.2026-27)"),
    ("Enterprises NOIDA", "ORANGE O TEC ENTERPRISES PRIVATE LIMITED-NOIDA -FY 26-27"),
    ("Colorix", "COLORIX DIGITAL PRINTING SOLUTIONS LLP - (from 1-Apr-20)"),
]
FAMILY_DEFS = [
    ("Total of Orange O Tec", "Total of Orange O Tec"),
    ("Total of Enterprises", "Total of Enterprises"),
    ("Grand Total", "Grand Total"),
]
HEADLINE_LABELS = [("Turnover", "TOTAL TURNOVER"), ("Gross Profit", "GROSS PROFIT"),
                    ("PAT (Net Profit)", "PAT (NET PROFIT)")]

NAVY = "1F3864"
GREEN = "C6E0B4"
RED = "F8CBAD"
YELLOW = "FFF2CC"
BLUE = "DDEBF7"
LIGHT_BLUE = "D9E2F3"
TITLE_FONT = Font(bold=True, size=14, name="Calibri", color=NAVY)
SUBTITLE_FONT = Font(italic=True, size=10, name="Calibri", color="595959")
SECTION_FONT = Font(bold=True, size=11, name="Calibri", color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, size=10, name="Calibri", color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2E5395")
NORMAL_FONT = Font(size=10, name="Calibri")
BOLD_FONT = Font(bold=True, size=10, name="Calibri")
TOTAL_FILL = PatternFill("solid", fgColor=LIGHT_BLUE)
NUMFMT = "#,##0"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
STATUS_FILL = {"ADDED": PatternFill("solid", fgColor=GREEN), "REMOVED": PatternFill("solid", fgColor=RED),
               "CHANGED": PatternFill("solid", fgColor=YELLOW), "RENAMED": PatternFill("solid", fgColor=BLUE)}


def norm(s):
    return re.sub(r"[^A-Z0-9]", "", s.upper())


def find_col(ws, header_row, text):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if isinstance(v, str) and v.strip() == text:
            return c
    raise ValueError(f"Column header {text!r} not found on row {header_row} of sheet {ws.title!r}")


def find_header_row(ws):
    for r in range(1, min(ws.max_row, 20) + 1):
        if ws.cell(r, 1).value == "Statement Section":
            return r
    raise ValueError(f"Header row ('Statement Section') not found in sheet {ws.title!r}")


def find_label_row(ws, label):
    up = label.strip().upper()
    for r in range(5, ws.max_row + 1):
        for c in (1, 4):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().upper() == up:
                return r
    raise ValueError(f"Row labelled {label!r} not found in sheet {ws.title!r}")


def collect_gl_values(ws, header_row, col):
    """{GL Name: summed literal value} for leaf rows only (skips Subtotal/TOTAL/header rows)."""
    d, ctx = {}, {}
    for r in range(header_row + 1, ws.max_row + 1):
        name = ws.cell(r, 4).value
        if not name or not isinstance(name, str):
            continue
        if name.startswith("Subtotal") or name.startswith("TOTAL") or "not booked" in name:
            continue
        val = ws.cell(r, col).value
        if isinstance(val, (int, float)):
            d[name] = d.get(name, 0) + val
            if name not in ctx:
                ctx[name] = (ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value)
    return d, ctx


def diff_gl_maps(omap, nmap, octx, nctx, threshold=1):
    all_names = set(omap) | set(nmap)
    added, removed = [], []
    rows = []
    for name in all_names:
        vo, vn = omap.get(name), nmap.get(name)
        vo_, vn_ = vo or 0, vn or 0
        delta = vn_ - vo_
        ctx = nctx.get(name) or octx.get(name) or (None, None, None)
        if vo is None and vn is not None:
            added.append({"name": name, "old": 0, "new": vn_, "delta": delta, "ctx": ctx})
        elif vn is None and vo is not None:
            removed.append({"name": name, "old": vo_, "new": 0, "delta": delta, "ctx": ctx})
        elif abs(delta) > threshold:
            rows.append({"name": name, "old": vo_, "new": vn_, "delta": delta, "ctx": ctx, "status": "CHANGED"})
    used = set()
    for a in added:
        akey = norm(a["name"])
        for i, r in enumerate(removed):
            if i in used:
                continue
            if norm(r["name"]) == akey:
                rows.append({"name": f'{r["name"]} -> {a["name"]}', "old": r["old"], "new": a["new"],
                             "delta": a["new"] - r["old"], "ctx": a["ctx"], "status": "RENAMED"})
                used.add(i)
                a["_matched"] = True
                break
    for a in added:
        if not a.get("_matched"):
            rows.append({**a, "status": "ADDED"})
    for i, r in enumerate(removed):
        if i not in used:
            rows.append({**r, "status": "REMOVED"})
    rows.sort(key=lambda x: -abs(x["delta"]))
    return rows


def compute_gl_diffs(old_wb, new_wb, sheetname):
    ow, nw = old_wb[sheetname], new_wb[sheetname]
    ohr, nhr = find_header_row(ow), find_header_row(nw)
    result = {}
    for cname, header_text in COMPANY_DEFS:
        ocol = find_col(ow, ohr, header_text)
        ncol = find_col(nw, nhr, header_text)
        omap, octx = collect_gl_values(ow, ohr, ocol)
        nmap, nctx = collect_gl_values(nw, nhr, ncol)
        result[cname] = diff_gl_maps(omap, nmap, octx, nctx)
    return result


def headline_values(wb, sheetname, label):
    ws = wb[sheetname]
    hr = find_header_row(ws)
    row = find_label_row(ws, label)
    out = {}
    for cname, header_text in COMPANY_DEFS + FAMILY_DEFS:
        col = find_col(ws, hr, header_text)
        out[cname] = ws.cell(row, col).value
    return out


def calc_headline(path, sheetname):
    """Evaluate the workbook with the `formulas` engine and pull headline values by label -
    never trust openpyxl's cached values (see SKILL.md 'Known bugs found & fixed' item 3)."""
    import formulas
    xl = formulas.ExcelModel().loads(str(path)).finish()
    sol = xl.calculate()
    # `formulas` key casing/quoting varies by workbook - build a case-insensitive lookup once
    # rather than a direct dict hit (a direct sol.get(key) silently returns None -> looks like 0).
    sol_upper = {k.upper(): v for k, v in sol.items()}
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb[sheetname]
    hr = find_header_row(ws)
    fname_key = Path(path).name.upper()

    def get(row, col):
        key = f"'[{fname_key}]{sheetname.upper()}'!{get_column_letter(col)}{row}"
        v = sol_upper.get(key.upper())
        if v is None:
            return None
        val = v.value
        try:
            return float(val[0][0])
        except Exception:
            return val

    out = {}
    for label, search_label in HEADLINE_LABELS:
        row = find_label_row(ws, search_label)
        out[label] = {}
        for cname, header_text in COMPANY_DEFS + FAMILY_DEFS:
            col = find_col(ws, hr, header_text)
            out[label][cname] = get(row, col)
    return out


def build_report(old_path, new_path, out_path):
    old_wb = openpyxl.load_workbook(old_path, data_only=False)
    new_wb = openpyxl.load_workbook(new_path, data_only=False)

    print("Calculating headline figures (old file) via formulas engine...")
    old_headline = calc_headline(old_path, "PL GL Detail")
    print("Calculating headline figures (new file) via formulas engine...")
    new_headline = calc_headline(new_path, "PL GL Detail")

    print("Computing GL-level diffs (P&L)...")
    pl_diff = compute_gl_diffs(old_wb, new_wb, "PL GL Detail")
    print("Computing GL-level diffs (Balance Sheet)...")
    bs_diff = compute_gl_diffs(old_wb, new_wb, "BS GL Detail")

    wb = openpyxl.Workbook()

    # ---------------- Sheet 1: PAT / headline reconciliation ----------------
    ws = wb.active
    ws.title = "PAT Reconciliation"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = f"Profit Reconciliation - {Path(old_path).stem} vs {Path(new_path).stem}"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A2:F2")
    ws["A2"] = ("Turnover / Gross Profit / PAT bridged between the previous snapshot and this "
                "refresh, per company. PAT drivers below are the actual top GL-level movements "
                "for whichever company/companies moved, auto-identified from the GL diff - not "
                "hand-picked.")
    ws["A2"].font = SUBTITLE_FONT
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 28
    row = 4

    for label, _ in HEADLINE_LABELS:
        ws.merge_cells(f"A{row}:F{row}")
        ws.cell(row, 1, value=label.upper()).font = SECTION_FONT
        for c in range(1, 7):
            ws.cell(row, c).fill = SECTION_FILL
        row += 1
        headers = ["Company / Family", "Old (Rs.)", "New (Rs.)", "Delta (Rs.)", "", ""]
        for c, h in enumerate(["Company / Family", "Old (Rs.)", "New (Rs.)", "Delta (Rs.)"], start=1):
            cell = ws.cell(row, c, value=h)
            cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
        row += 1
        for cname, _ in COMPANY_DEFS + FAMILY_DEFS:
            is_total = cname == "Grand Total"
            vo = old_headline[label].get(cname) or 0
            vn = new_headline[label].get(cname) or 0
            d = vn - vo
            for c, v in enumerate([cname, vo, vn, d], start=1):
                cell = ws.cell(row, c, value=v)
                cell.border = BORDER
                cell.font = BOLD_FONT if is_total else NORMAL_FONT
                if is_total:
                    cell.fill = TOTAL_FILL
                if c > 1:
                    cell.number_format = NUMFMT
            row += 1

        if label == "PAT (Net Profit)":
            row += 1
            ws.merge_cells(f"A{row}:F{row}")
            ws.cell(row, 1, value="TOP GL-LEVEL DRIVERS BY COMPANY (auto-identified)").font = SECTION_FONT
            for c in range(1, 7):
                ws.cell(row, c).fill = SECTION_FILL
            row += 1
            headers2 = ["Company", "GL Name", "Old (Rs.)", "New (Rs.)", "Delta (Rs.)", ""]
            for c, h in enumerate(["Company", "GL Name", "Old (Rs.)", "New (Rs.)", "Delta (Rs.)"], start=1):
                cell = ws.cell(row, c, value=h)
                cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
            row += 1
            company_deltas = []
            for cname, _ in COMPANY_DEFS:
                vo = old_headline[label].get(cname) or 0
                vn = new_headline[label].get(cname) or 0
                company_deltas.append((cname, vn - vo))
            company_deltas.sort(key=lambda x: -abs(x[1]))
            for cname, cdelta in company_deltas:
                if abs(cdelta) < 10000:
                    continue
                drivers = pl_diff.get(cname, [])
                running = 0.0
                shown = 0
                for d in drivers:
                    if shown >= 5 or (running / cdelta if cdelta else 1) >= 0.95:
                        break
                    vals = [cname if shown == 0 else "", d["name"], d["old"], d["new"], d["delta"]]
                    for c, v in enumerate(vals, start=1):
                        cell = ws.cell(row, c, value=v)
                        cell.border = BORDER
                        cell.font = NORMAL_FONT
                        if c >= 3:
                            cell.number_format = NUMFMT
                        if c == 2:
                            cell.alignment = Alignment(wrap_text=True, vertical="top")
                    row += 1
                    running += d["delta"]
                    shown += 1
        row += 2

    widths = [26, 45, 16, 16, 16, 4]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    # ---------------- Sheets 2 & 3: GL Changes ----------------
    def build_gl_change_sheet(sheetname, title, diff_data):
        ws = wb.create_sheet(sheetname)
        ws.sheet_view.showGridLines = False
        ws.merge_cells("A1:H1")
        ws["A1"] = title
        ws["A1"].font = TITLE_FONT
        ws.merge_cells("A2:H2")
        ws["A2"] = ("Every GL ledger Added, Removed, Changed in value, or Renamed (same ledger, "
                    "different Tally text) between the two files. Sorted by company, then size of "
                    "impact. Only real movements >Rs.1 shown.")
        ws["A2"].font = SUBTITLE_FONT
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 28
        row = 4
        headers = ["Company", "Status", "Statement Section", "Grouping", "Sub Grouping",
                   "GL Name (Old -> New if renamed)", "Old (Rs.)", "New (Rs.)"]
        headers.append("Delta (Rs.)")
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row, c, value=h)
            cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 26
        row += 1
        header_row = row - 1
        for cname, rows_ in diff_data.items():
            if not rows_:
                continue
            added = sum(1 for r in rows_ if r["status"] == "ADDED")
            removed = sum(1 for r in rows_ if r["status"] == "REMOVED")
            changed = sum(1 for r in rows_ if r["status"] == "CHANGED")
            renamed = sum(1 for r in rows_ if r["status"] == "RENAMED")
            net = sum(r["delta"] for r in rows_)
            ws.merge_cells(f"A{row}:I{row}")
            ws.cell(row, 1, value=(f"{cname}  -  {added} added, {removed} removed, {changed} changed, "
                                    f"{renamed} renamed  |  net impact Rs. {net:,.0f}")).font = SECTION_FONT
            for c in range(1, 10):
                ws.cell(row, c).fill = SECTION_FILL
            row += 1
            for r in rows_:
                sec, grp, subgrp = r["ctx"]
                vals = [None, r["status"], sec, grp, subgrp, r["name"], r["old"], r["new"], r["delta"]]
                for c, v in enumerate(vals, start=1):
                    cell = ws.cell(row, c, value=v)
                    cell.border = BORDER
                    cell.font = NORMAL_FONT
                    if c == 2:
                        cell.fill = STATUS_FILL.get(r["status"], PatternFill())
                        cell.alignment = Alignment(horizontal="center")
                    if c >= 7:
                        cell.number_format = NUMFMT
                    if c == 6:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                row += 1
        widths = [30, 11, 16, 22, 24, 48, 16, 16, 16]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:I{header_row}"

    build_gl_change_sheet("GL Changes - PL", "GL-Level Changes - Profit & Loss (all 5 companies)", pl_diff)
    build_gl_change_sheet("GL Changes - BS", "GL-Level Changes - Balance Sheet (all 5 companies)", bs_diff)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print("Saved:", out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--old", default=None, help="Defaults to the 2nd-newest dated file in Files/.")
    ap.add_argument("--new", default=None, help="Defaults to the newest dated file in Files/.")
    ap.add_argument("--out", default=None, help="Defaults to Profit_Change_Reconciliation_<new's timestamp>.xlsx")
    args = ap.parse_args()

    if args.old and args.new:
        old_path, new_path = Path(args.old), Path(args.new)
    else:
        candidates = find_pl_bs_files(FILES_DIR)
        if len(candidates) < 2:
            print(f"Found {len(candidates)} dated PL_BS_GL_Detail_*.xlsx file(s) in {FILES_DIR} - "
                  f"need at least 2 (a previous one and this refresh's) to produce a reconciliation "
                  f"report. Skipping this run; it will be available starting from the next refresh.")
            sys.exit(0)
        old_path, new_path = candidates[-2], candidates[-1]

    if not old_path.exists():
        print(f"ERROR: old file not found: {old_path}")
        sys.exit(1)
    if not new_path.exists():
        print(f"ERROR: new file not found: {new_path}")
        sys.exit(1)

    if args.out:
        out_path = Path(args.out)
    else:
        ts = extract_timestamp(new_path) or datetime.now().strftime(TIMESTAMP_FORMAT)
        out_path = FILES_DIR / f"{RECON_BASE_NAME}_{ts}.xlsx"

    build_report(old_path, new_path, out_path)


if __name__ == "__main__":
    main()
