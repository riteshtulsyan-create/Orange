"""
publish_cfo_pack_hub.py - put a CFO Pack into Orange Hub (local server) for viewing and download.

    python publish_cfo_pack_hub.py                      # newest Files\\Results\\CFO_Pack_*.xlsx
    python publish_cfo_pack_hub.py --pack <xlsx> --hub "<Orange Hub folder>"

Copies the pack's .xlsx and .pdf into <hub>\\frontend\\public\\reports\\cfo-pack\\ and writes
summary.json there - the summary bullets and the tables of the 'CFO Analysis' sheet, read from the
values Excel saved when it built the pack. The Hub page /reports/finance-cfo-pack reads that JSON.

That folder is GITIGNORED in the Hub on purpose: company financials stay on this machine's server
and never travel to GitHub with the code.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT = SCRIPT_DIR.parent.parent.parent.parent
RESULTS = PROJECT / "Files" / "Results"
DEFAULT_HUB = PROJECT.parent / "Orange Hub - Yash Agarwal" / "Orange Hub-Ritesh-PL"

SECTION = re.compile(r"^(\d+b?)\.\s|^[A-D]\.\s|^SUMMARY FOR THE CFO|^Annex [A-Z] - ")


def read_analysis(xlsx: Path, sheet: str = "CFO Analysis"):
    tmp = Path(tempfile.mkdtemp(prefix="cfo_hub_")) / xlsx.name
    shutil.copy2(xlsx, tmp)
    ws = openpyxl.load_workbook(tmp, data_only=True, read_only=True)[sheet]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if sheet == "CFO Analysis" and not any(isinstance(v, (int, float)) for r in rows for v in r[1:]):
        sys.exit(f"{xlsx.name} has no calculated values - it must be the file Excel saved (build without --no-pdf)")
    sections, cur = [], None
    for r in rows[3:]:
        a = r[0]
        cells = [v for v in r if v not in (None, "")]
        if not cells:
            continue
        if isinstance(a, str) and SECTION.match(a.strip()):
            cur = {"title": a.strip(), "header": None, "rows": [], "notes": [], "bullets": []}
            sections.append(cur)
            continue
        if cur is None:
            continue
        if isinstance(a, str) and a.startswith("•"):
            cur["bullets"].append({"group": cur.get("_group"), "text": a.lstrip("• ").strip()})
            continue
        nums = [v for v in r[1:] if isinstance(v, (int, float))]
        texts = [v for v in r[1:] if isinstance(v, str) and v.strip()]
        if not nums and not texts and isinstance(a, str):
            if re.match(r"^[A-D]\.\s", a):
                cur["_group"] = a.strip()
            else:
                cur["notes"].append(a.strip())
            continue
        if cur["header"] is None and not nums:
            last = max(i for i, v in enumerate(r) if v not in (None, ""))
            cur["header"] = [str(v).replace("\n", " ") if v is not None else "" for v in r[:last + 1]]
            continue
        width = len(cur["header"] or r)
        cur["rows"].append([v if not isinstance(v, float) else round(v, 4) for v in r[:width]])
    for s in sections:
        s.pop("_group", None)
    # Annexes are working controls, grouped and hidden in the Excel - not part of the report (user, 2026-10-24).
    return [s for s in sections if not s["title"].startswith("Annex")]


def read_mom(xlsx: Path):
    """The 'PDF - MOM' sheet: one block per view (Group / O Tec / Enterprises / Colorix)."""
    tmp = Path(tempfile.mkdtemp(prefix="cfo_hub_")) / xlsx.name
    shutil.copy2(xlsx, tmp)
    ws = openpyxl.load_workbook(tmp, data_only=True, read_only=True)["PDF - MOM"]
    views, cur = [], None
    for r in ws.iter_rows(min_row=4, values_only=True):
        a = r[0]
        if isinstance(a, str) and a.startswith("MONTH ON MONTH - "):
            cur = {"name": a[len("MONTH ON MONTH - "):].strip(), "columns": None, "rows": []}
            views.append(cur)
        elif cur is not None and a == "Rs crore":
            cur["columns"] = [v for v in r[1:] if v]
        elif cur is not None and isinstance(a, str) and a.strip():
            n = len(cur["columns"] or [])
            vals = [round(v, 4) if isinstance(v, (int, float)) else None for v in r[1:1 + n]]
            cur["rows"].append({"label": a.strip(), "band": all(v is None for v in vals), "values": vals})
    return views


def load_detail(path: Path):
    """The builder's sidecar: {'mom_detail': ..., 'memo': {...}} (older packs: the drill-down alone)."""
    if not path.exists():
        return {"mom_detail": {}, "memo": {}}
    d = json.loads(path.read_text(encoding="utf-8"))
    return d if "mom_detail" in d else {"mom_detail": d, "memo": {}}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pack", help="CFO_Pack_*.xlsx (default: newest in Files\\Results)")
    ap.add_argument("--hub", help=f"Orange Hub folder (default: {DEFAULT_HUB})")
    args = ap.parse_args()
    # The crore, detailed pack is the main one (the Hub data is read from it); its lakh / absolute and
    # short versions share its stamp.
    main_pack = re.compile(r"CFO_Pack_PL_BS_CF_FF_\d{4}-\d\d-\d\d_\d{4}\.xlsx$")
    pack = Path(args.pack) if args.pack else max((p for p in RESULTS.glob("CFO_Pack_PL_BS_CF_FF_*.xlsx")
                                                  if main_pack.match(p.name)), key=lambda p: p.name)
    pdf = pack.with_suffix(".pdf")
    hub = Path(args.hub) if args.hub else DEFAULT_HUB
    dest = hub / "frontend" / "public" / "reports" / "cfo-pack"
    if not (hub / "frontend").exists():
        sys.exit(f"Not an Orange Hub folder: {hub}")
    dest.mkdir(parents=True, exist_ok=True)
    detail = pack.with_name(pack.stem + "_detail.json")
    stamp = pack.stem.replace("CFO_Pack_PL_BS_CF_FF_", "")
    dpdf = pack.with_name(f"Directors_Executive_Summary_{stamp}.pdf")
    for pattern in ("CFO_Pack_*", "Directors_Executive_Summary_*"):
        for old in dest.glob(pattern):
            old.unlink()
    # Every download, by unit and length: {unit: {detailed: {xlsx, pdf}, short: {xlsx, pdf}, directors}}.
    downloads = {}
    for unit, sfx in (("cr", ""), ("lakh", "_lakh"), ("abs", "_abs")):
        files = {
            "detailed": {"xlsx": f"CFO_Pack_PL_BS_CF_FF_{stamp}{sfx}.xlsx", "pdf": f"CFO_Pack_PL_BS_CF_FF_{stamp}{sfx}.pdf"},
            "short": {"xlsx": f"CFO_Pack_PL_BS_CF_FF_{stamp}{sfx}_short.xlsx", "pdf": f"CFO_Pack_PL_BS_CF_FF_{stamp}{sfx}_short.pdf"},
            "directors": f"Directors_Executive_Summary_{stamp}{sfx}.pdf",
        }
        names = [files["directors"], *files["detailed"].values(), *files["short"].values()]
        if not all((pack.parent / n).exists() for n in names):
            continue
        for n in names:
            shutil.copy2(pack.parent / n, dest / n)
        downloads[unit] = files
    for f in (pack, pdf, dpdf):
        if f.exists() and not (dest / f.name).exists():
            shutil.copy2(f, dest / f.name)
    summary = {
        "built": dt.datetime.fromtimestamp(pack.stat().st_mtime).strftime("%d-%b-%Y %H:%M"),
        "xlsx": pack.name,
        "pdf": pdf.name if pdf.exists() else None,
        "directors_pdf": dpdf.name if dpdf.exists() else None,
        "downloads": downloads,
        "sections": read_analysis(pack),
        "mom": read_mom(pack),
        # {view: {line: [{sub, values, ledgers: [{gl, values}]}]}} - the GL drop-down
        **load_detail(detail),
        "briefing": read_analysis(pack, "CFO Briefing"),
        "directors": read_analysis(pack, "Directors Summary"),
    }
    (dest / "summary.json").write_text(json.dumps(summary, indent=1, default=str), encoding="utf-8")
    print(f"Published to {dest}: {pack.name}, {pdf.name if pdf.exists() else '(no pdf)'}, summary.json "
          f"({len(summary['sections'])} sections); download versions: {', '.join(downloads) or 'none'}")


if __name__ == "__main__":
    main()
