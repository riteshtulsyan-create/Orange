"""
Fetch Profit & Loss and Balance Sheet from Tally (port 9010) for all open companies,
laid out side-by-side (one amount column per company) matching the structure of
Files/PL FOMRAT.xlsx.

Usage:
    python fetch_pandl_bs_combined.py --from-date 01-04-2026 --to-date 30-06-2026 --output "Files/PL_BS_Combined_Output.xlsx"
"""

import argparse
import os
import sys
import urllib.request
import urllib.error
import xml.etree.ElementTree as ET
from datetime import datetime, date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

# Load .env: check CWD first (project root), then script dir, then parent
_env_candidates = [
    Path.cwd() / ".env",
    Path(__file__).parent / ".env",
    Path(__file__).parent.parent / ".env",
]
for _env_path in _env_candidates:
    if _env_path.exists():
        for _line in _env_path.read_text().splitlines():
            _line = _line.strip()
            if _line and not _line.startswith("#") and "=" in _line:
                _k, _v = _line.split("=", 1)
                os.environ.setdefault(_k.strip(), _v.strip())
        break

TALLY_HOST = os.environ.get("TALLY_HOST", "localhost")
TALLY_PORT = int(os.environ.get("TALLY_PORT", "9010"))

# ---------------------------------------------------------------------------
# P&L line-item order (from Files/PL FOMRAT.xlsx, "PL YTD" sheet)
# ---------------------------------------------------------------------------
PL_ACCOUNT_ORDER = [
    ("Trading Account:", "group_c", "C"),
    ("Sales Accounts", "group_c", "C"),
    ("DISCOUNT AND RATE DIFFERENCE (GST)", "ledger", "B"),
    ("CONSUMPTION CHARGES (INCOME) 18%", "ledger", "B"),
    ("DISCOUNT & RATE DIFFERENCE@12% (PAPER)", "ledger", "B"),
    ("DISCOUNT & RATE DIFFERENCE@18%-HEAD", "ledger", "B"),
    ("DISCOUNT & RATE DIFFERENCE@18% (INK)", "ledger", "B"),
    ("DISCOUNT & RATE DIFFERENCE@18% (MACHINE)", "ledger", "B"),
    ("DISCOUNT & RATE DIFFERENCE@18% (PAPER)", "ledger", "B"),
    ("DISCOUNT & RATE DIFFERENCE@18%-SPARE", "ledger", "B"),
    ("RATE DIFFERENCE (INCOME) 18%-HEAD", "ledger", "B"),
    ("RATE DIFFERENCE (INCOME) 18% (INK)", "ledger", "B"),
    ("RATE DIFFERENCE (INCOME) 18%-MACHINE", "ledger", "B"),
    ("RATE DIFFERENCE (INCOME) 18%-PAPER", "ledger", "B"),
    ("RATE DIFFERENCE (INCOME) 18%-SPARE", "ledger", "B"),
    ("DISCOUNT AND RATE DIFFERENCE (NON GST)", "ledger", "B"),
    ("DISCOUNT", "ledger", "B"),
    ("DISCOUNT (BD)", "ledger", "B"),
    ("DISCOUNT ON SALE(NON-GST)", "ledger", "B"),
    ("DISCOUNT & RATE DIFFERENCE (NON GST)-MACHINE", "ledger", "B"),
    ("RELATED PARTY SALES", "ledger", "B"),
    ("RELATED PARTY SALES-HEAD", "ledger", "B"),
    ("RELATED PARTY SALES-INK", "ledger", "B"),
    ("RELATED PARTY SALES-MACHINE", "ledger", "B"),
    ("RELATED PARTY SALES-PACKING MATERIAL", "ledger", "B"),
    ("RELATED PARTY SALES-PAPER", "ledger", "B"),
    ("RELATED PARTY SALES RETURN-HEAD", "ledger", "B"),
    ("RELATED PARTY SALES RETURN-INK", "ledger", "B"),
    ("RELATED PARTY SALES RETURN-SPARE", "ledger", "B"),
    ("RELATED PARTY SALES-SPARE", "ledger", "B"),
    ("GST BRANCH SALE ACCOUNTS", "ledger", "B"),
    ("GST INTERSTATE SALE - HEAD", "ledger", "B"),
    ("GST INTERSTATE SALE -INK", "ledger", "B"),
    ("GST INTERSTATE SALE- MACHINE", "ledger", "B"),
    ("GST INTERSTATE SALE-PAPER", "ledger", "B"),
    ("GST INTERSTATE SALE-SPARE", "ledger", "B"),
    ("GST SALES ACCOUNTS -HEAD", "ledger", "B"),
    ("GST SALES ACCOUNTS - INK", "ledger", "B"),
    ("GST SALES ACCOUNTS - MACHINE", "ledger", "B"),
    ("GST SALES ACCOUNTS-PAPER", "ledger", "B"),
    ("GST SALES ACCOUNTS - SPARE", "ledger", "B"),
    ("GST SALES RETURN-HEAD", "ledger", "B"),
    ("GST SALES RETURN-INK", "ledger", "B"),
    ("GST SALES RETURN-MACHINE", "ledger", "B"),
    ("GST SALES RETURN-PAPER", "ledger", "B"),
    ("GST SALES RETURN-SPARE", "ledger", "B"),
    ("GST SALES RETURN-SPARE_NON GST", "ledger", "B"),
    ("HIGH SEAS SALES", "ledger", "B"),
    ("SALES EXPORT", "ledger", "B"),
    ("Sales Bills to Make", "group_b", "B"),
    ("Direct Incomes", "group_c", "C"),
    ("COURIER INCOME", "ledger", "B"),
    ("COURIER CHARGES 5%", "ledger", "B"),
    ("COURIER CHARGES (INC)", "ledger", "B"),
    ("FREIGHT INCOME", "ledger", "B"),
    ("CLEARING & FORWERDING CHARGES CENTRING (INCOME)-18%", "ledger", "B"),
    ("CLEARING & FORWERDING CHARGES HEAD- (INCOME)-18%", "ledger", "B"),
    ("CLEARING & FORWERDING CHARGES INK- (INCOME)-18%", "ledger", "B"),
    ("CLEARING & FORWERDING CHARGES MACHINE- (INCOME)-18%", "ledger", "B"),
    ("FREIGHT CHARGES@18% (INC)", "ledger", "B"),
    ("TRANSPORT CHARGES (INC)", "ledger", "B"),
    ("OTHER INCOME", "ledger", "B"),
    ("ANNUAL MAINTENANCE CONTRACT (INC)", "ledger", "B"),
    ("HANDLING CHARGES ACCOUNTS (INC)", "ledger", "B"),
    ("SERVICE CHARGES ACCOUNTS (INC)", "ledger", "B"),
    ("", "subtotal_c", "C"),
    ("Cost of Sales :", "group_c", "C"),
    ("Opening Stock", "ledger", "B"),
    ("Add: Purchase Accounts", "ledger", "B"),
    ("Less: Closing Stock", "ledger", "B"),
    ("", "subtotal_b", "B"),
    ("Direct Expenses", "group_b", "B"),
    ("CARRIAGE INWARD EXP", "ledger", "B"),
    ("CARRIAGE INWARD ACCOUNTS - 12%", "ledger", "B"),
    ("CARRIAGE INWARD EXP ACCOUNTS-18%", "ledger", "B"),
    ("CARRIAGE INWARD EXP ACCOUNTS@5%", "ledger", "B"),
    ("CARRIAGE INWARD EXP ACCOUNTS-NON GST", "ledger", "B"),
    ("CARRIAGE INWARD EXP ACCOUNTS (RCM)-FABRIC", "ledger", "B"),
    ("CARRIAGE INWARD EXP ACCOUNTS (RCM)-HEAD", "ledger", "B"),
    ("CARRIAGE INWARD EXP ACCOUNTS (RCM)-INK", "ledger", "B"),
    ("CARRIAGE INWARD EXP ACCOUNTS (RCM)-MACHINE", "ledger", "B"),
    ("CARRIAGE INWARD EXP ACCOUNTS (RCM)-PAPER", "ledger", "B"),
    ("CARRIAGE INWARD EXP ACCOUNTS (RCM)-POWDER", "ledger", "B"),
    ("CARRIAGE INWARD EXP ACCOUNTS (RCM)-SPARE", "ledger", "B"),
    ("CLEARING & FORWERDING EXP", "ledger", "B"),
    ("CLEARING AND TRANSPORT CHARGES", "ledger", "B"),
    ("CLEARING & FORWERDING CHARGES (NON-GST)-HEAD", "ledger", "B"),
    ("CLEARING & FORWERDING CHARGES (NON-GST)-INK", "ledger", "B"),
    ("CLEARING & FORWERDING CHARGES (NON-GST)-MACHINE", "ledger", "B"),
    ("CLEARING & FORWERDING CHARGES (NON-GST)-PAPER", "ledger", "B"),
    ("CLEARING & FORWERDING CHARGES (NON-GST)-POWDER", "ledger", "B"),
    ("CLEARING & FORWERDING CHARGES (NON-GST)-SPARE", "ledger", "B"),
    ("CLEARING & FORWERDING CHARGES-RCM", "ledger", "B"),
    ("CLEARNING & FORWERDING CHARGES (12%)-PAPER", "ledger", "B"),
    ("CLEARNING & FORWERDING CHARGES (5%)-PAPER", "ledger", "B"),
    ("CLEARNING & FORWERDING CHARGES (ITC)-FABRIC", "ledger", "B"),
    ("CLEARNING & FORWERDING CHARGES (ITC)-HEAD", "ledger", "B"),
    ("CLEARNING & FORWERDING CHARGES (ITC)-INK", "ledger", "B"),
    ("CLEARNING & FORWERDING CHARGES (ITC)-MACHINE", "ledger", "B"),
    ("CLEARNING & FORWERDING CHARGES (ITC)-MACHINE-EXPORT", "ledger", "B"),
    ("CLEARNING & FORWERDING CHARGES (ITC)-PAPER", "ledger", "B"),
    ("CLEARNING & FORWERDING CHARGES (ITC)-POWDER", "ledger", "B"),
    ("CLEARNING & FORWERDING CHARGES (ITC)-SPARE", "ledger", "B"),
    ("CUSTOM DUTY EXPENSES", "ledger", "B"),
    ("CUSTOM DUTY EXP ACCOUNTS-FABRICS", "ledger", "B"),
    ("CUSTOM DUTY EXP ACCOUNTS-HEAD", "ledger", "B"),
    ("CUSTOM DUTY EXP ACCOUNTS-INK", "ledger", "B"),
    ("CUSTOM DUTY EXP ACCOUNTS-MACHINE", "ledger", "B"),
    ("CUSTOM DUTY EXP ACCOUNTS-PAPER", "ledger", "B"),
    ("CUSTOM DUTY EXP ACCOUNTS-SPARE", "ledger", "B"),
    ("INTEREST ON DUTY", "ledger", "B"),
    ("INTEREST ON DUTY-PAPER", "ledger", "B"),
    ("EXCHANGE RATE Gain/(Loss)", "ledger", "B"),
    ("EXCHANGE RATE DIFF EXP ACCOUNTS (PURCHASE)", "ledger", "B"),
    ("EXCHANGE RATE DIFF EXP ACCOUNTS (PURCHASE)-OLD", "ledger", "B"),
    ("EXCHANGE RATE DIFF EXP ACCOUNTS (RECEIPT)", "ledger", "B"),
    ("EXCHANGE RATE DIFF EXP ACCOUNTS_LC-BC", "ledger", "B"),
    ("EXCHANGE RATE DIFF EXP AC-MAR-26", "ledger", "B"),
    ("FREIGHT EXP", "ledger", "B"),
    ("AIR FREIGHT", "ledger", "B"),
    ("AIR FREIGHT CHARGES @ 5%-SPARE", "ledger", "B"),
    ("AIR FREIGHT CHARGES (NON-GST)-HEAD", "ledger", "B"),
    ("AIR FREIGHT CHARGES (NON-GST)-INK", "ledger", "B"),
    ("AIR FREIGHT CHARGES (NON-GST)-PAPER", "ledger", "B"),
    ("AIR FREIGHT CHARGES (NON-GST)-SPARE", "ledger", "B"),
    ("AIR FREIGHT EXP ACCOUNTS@18%", "ledger", "B"),
    ("AIR FREIGHT EXP ACCOUNTS@18%-HEAD", "ledger", "B"),
    ("AIR FREIGHT EXP ACCOUNTS@18% (INK)", "ledger", "B"),
    ("AIR FREIGHT EXP ACCOUNTS@18%-PAPER", "ledger", "B"),
    ("AIR FREIGHT EXP ACCOUNTS@18%-SPARE", "ledger", "B"),
    ("AIR FREIGHT EXP ACCOUNTS(INK)- Export", "ledger", "B"),
    ("OCEAN FREIGHT", "ledger", "B"),
    ("OCEAN FREIGHT CHARGES @ 18% (HEAD)", "ledger", "B"),
    ("OCEAN FREIGHT CHARGES @ 5%-(INK)", "ledger", "B"),
    ("OCEAN FREIGHT CHARGES @ 5% (MACHINE)", "ledger", "B"),
    ("OCEAN FREIGHT CHARGES @ 5%-(PAPER)", "ledger", "B"),
    ("OCEAN FREIGHT CHARGES @ 5% (POWDER)", "ledger", "B"),
    ("OCEAN FREIGHT CHARGES @ 5% (SPARE)", "ledger", "B"),
    ("OCEAN FREIGHT CHARGES (MACHINE)18%- Export", "ledger", "B"),
    ("OCEAN FREIGHT CHARGES (MACHINE)5%- Export", "ledger", "B"),
    ("OCEAN FREIGHT CHARGES (MACHINE)- Export", "ledger", "B"),
    ("OCEAN FREIGHT EXP ACCOUNTS@12%-MACHINE", "ledger", "B"),
    ("OCEAN FREIGHT EXP ACCOUNTS@18% (INK)", "ledger", "B"),
    ("OCEAN FREIGHT EXP ACCOUNTS@18% (INK)-EXPORT", "ledger", "B"),
    ("OCEAN FREIGHT EXP ACCOUNTS@18% (MACHINE)", "ledger", "B"),
    ("OCEAN FREIGHT EXP ACCOUNTS@18%(PAPER)", "ledger", "B"),
    ("OCEAN FREIGHT EXP ACCOUNTS@18% (POWDER)", "ledger", "B"),
    ("OCEAN FREIGHT EXP ACCOUNTS@18% (SPARE)", "ledger", "B"),
    ("OCEAN FREIGHT EXP ACCOUNTS (NON GST)-INK", "ledger", "B"),
    ("OCEAN FREIGHT EXP ACCOUNTS (NON GST)-MACHINE", "ledger", "B"),
    ("OCEAN FREIGHT EXP ACCOUNTS (NON GST)-(PAPER)", "ledger", "B"),
    ("TRANSPORT CHARGES  (PURCHASE)", "ledger", "B"),
    ("FREIGHT CHARGE- (12%)", "ledger", "B"),
    ("FREIGHT CHARGES (18%)", "ledger", "B"),
    ("FREIGHT CHARGES-NON GST", "ledger", "B"),
    ("JOB WORK CHARGES", "ledger", "B"),
    ("JOB WORK EXP@18% (Related)", "ledger", "B"),
    ("JOB WORK EXP@5%", "ledger", "B"),
    ("Other Diretc Expenses", "ledger", "B"),
    ("AMC CHARGES EXP - 18%", "ledger", "B"),
    ("LOCAL SPARE PARTS EXP", "ledger", "B"),
    ("SERVICE CHARGES ACCOUNTS (MACHINERY & PARTS)", "ledger", "B"),
    ("SERVICE CHARGES ACCOUNTS (MACHINERY & PARTS)-18%", "ledger", "B"),
    ("PACKING EXP", "ledger", "B"),
    ("PACKAGING CHARGERS ON IMPORT", "ledger", "B"),
    ("PACKING MATERIAL CHAGRES", "ledger", "B"),
    ("PACKING MATRIAL ACCOUNTS", "ledger", "B"),
    ("PACKING MATRIAL ACCOUNTS 18%", "ledger", "B"),
    ("SAMPLE FABRIC EXPENSES", "ledger", "B"),
    ("SAMPLE FABRIC PURCHASE EXP.(18%)", "ledger", "B"),
    ("SAMPLE FABRIC PURCHASE EXP.(5%)", "ledger", "B"),
    ("SAMPLE FABRIC PURCHASE EXP(NON-GST)", "ledger", "B"),
    ("SERVICE EXP", "ledger", "B"),
    ("SERVICE CHARGE (NON-GST)", "ledger", "B"),
    ("SERVICE CHARGES ACCOUNTS (ITC)", "ledger", "B"),
    ("INVENTORY OVERHEAD COST", "ledger", "B"),
    ("INVENTORY - TRANSFER TO FA", "ledger", "B"),
    ("PROVISION FOR IMPORT-FOC(PL)", "ledger", "B"),
    ("Gross Profit :", "group_c", "C"),
    ("Income Statement:", "group_c", "C"),
    ("Indirect Incomes", "group_c", "C"),
    ("INTEREST INCOME", "ledger", "B"),
    ("INTEREST INCOME ON UNSECURED LOAN A/C(RELATED)", "ledger", "B"),
    ("FD INTEREST RECEIVED ACCOUNTS", "ledger", "B"),
    ("INTEREST ON FD", "ledger", "B"),
    ("OTHER INCOME", "ledger", "B"),
    ("PROFIT ON SALE OF ASSET", "ledger", "B"),
    ("RENT (INCOME)", "ledger", "B"),
    ("RENT (INCOME-RELATED)", "ledger", "B"),
    ("", "subtotal_c", "C"),
    ("Indirect Expenses", "group_c", "C"),
    ("ADVERTISING PROMOTION EXP", "ledger", "B"),
    ("ADVERTISING EXP ACCOUNTS @5%", "ledger", "B"),
    ("ADVERTISING EXP ACCOUNTS-NON GST", "ledger", "B"),
    ("ADVERTISING EXP(ITC) - 18%", "ledger", "B"),
    ("EXHIBITION EXP - 18%", "ledger", "B"),
    ("EXHIBITION EXP - 5%", "ledger", "B"),
    ("EXHIBITION EXP-NON GST", "ledger", "B"),
    ("COMMISSION & BROKERAGE EXP", "ledger", "B"),
    ("COMMISSION & BROKERAGE ACCOUNTS-E", "ledger", "B"),
    ("COMMISSION & BROKERAGE ACCOUNTS (ITC)", "ledger", "B"),
    ("COMMISSION & BROKERAGE ACCOUNTS-NON GST", "ledger", "B"),
    ("COMPUTER & PRINTER EXP", "ledger", "B"),
    ("COMPUTER & PRINTER EXPENSES", "ledger", "B"),
    ("COMPUTER & PRINTER EXPENSES 18%", "ledger", "B"),
    ("INTERNET & WEBSITE EXPENSES", "ledger", "B"),
    ("SOFTWARE CHARGES  ACCOUNTS", "ledger", "B"),
    ("TONER REFELING EXP-18%", "ledger", "B"),
    ("DEPRECIATION & AMORTISATION EXPENSE", "ledger", "B"),
    ("DEPRECIATION", "ledger", "B"),
    ("DONATION SUBSCRIPTION", "ledger", "B"),
    ("DONATION EXP", "ledger", "B"),
    ("DONATION EXP(CSR)", "ledger", "B"),
    ("ELECTRICITY EXPENSES", "ledger", "B"),
    ("ELECTRICITY EXPENSES ACCOUNTS", "ledger", "B"),
    ("EMPLOYEE COST", "ledger", "B"),
    ("EMPLOYEE BENEFITS EXPENSES", "ledger", "B"),
    ("EMPLOYEE LABOUR WELFARE CHARGES", "ledger", "B"),
    ("EMPLOYER ESIC CONTRIBUTION", "ledger", "B"),
    ("EMPLOYER LABOUR WELFARE CONTRIBUTION", "ledger", "B"),
    ("EMPLOYER PF CONTRIBUTION", "ledger", "B"),
    ("EMPLOYER PROFESSIONAL TAX", "ledger", "B"),
    ("INTEREST ON PT", "ledger", "B"),
    ("PENALTY CHARGES ON PF", "ledger", "B"),
    ("PENALTY CHARGES ON PT", "ledger", "B"),
    ("PF ADMIN CHARGES", "ledger", "B"),
    ("EMPLOYEE SALARY EXPENSES", "ledger", "B"),
    ("DRIVER SALARY EXP A/C", "ledger", "B"),
    ("EMPLOYEE BONUS ACCOUNTS", "ledger", "B"),
    ("EMPLOYEE INCENTIVES EXP ACCOUNTS", "ledger", "B"),
    ("EMPLOYEE PERFORMANCE EXPENSES", "ledger", "B"),
    ("EMPLOYEE SALARY ACCOUNTS", "ledger", "B"),
    ("GRATUITY EXPENSES", "ledger", "B"),
    ("MANAGERIAL REMUNERATION", "ledger", "B"),
    ("DIRECTORS SALARY EXP ACCOUNTS", "ledger", "B"),
    ("STAFF WELFARE EXPENSES", "ledger", "B"),
    ("BUSINESS PROMOTION EXPENSE- NON GST", "ledger", "B"),
    ("STAFF RECRUITMENT EXPESES-18%", "ledger", "B"),
    ("STAFF WELFARE EXP 18%", "ledger", "B"),
    ("STAFF WELFARE EXP 5%", "ledger", "B"),
    ("STAFF WELFARE EXP ACCOUNTS", "ledger", "B"),
    ("EMPLOYEE TRAVEL COST", "ledger", "B"),
    ("HOTEL EXPENSES", "ledger", "B"),
    ("HOTEL EXP ACCOUNT (12%)", "ledger", "B"),
    ("HOTEL EXP ACCOUNT (18%)", "ledger", "B"),
    ("HOTEL EXP ACCOUNT (5%)", "ledger", "B"),
    ("HOTEL EXP ACCOUNT (INTERNATIONAL)", "ledger", "B"),
    ("HOTEL EXP ACCOUNT (NON-GST)", "ledger", "B"),
    ("TRAVELLING EXP", "ledger", "B"),
    ("TC-D/A", "ledger", "B"),
    ("TC-FUEL", "ledger", "B"),
    ("TC-HOTEL", "ledger", "B"),
    ("TC-MEALS", "ledger", "B"),
    ("TC-OTHERS", "ledger", "B"),
    ("TC-TRAVEL", "ledger", "B"),
    ("TRAVELLING AND CON. EXP-INT.DIRECTOR", "ledger", "B"),
    ("TRAVELLING AND CON. EXP-INT.DIRECTOR-18%", "ledger", "B"),
    ("TRAVELLING AND CON. EXP-INTERNATIONAL", "ledger", "B"),
    ("TRAVELLING AND CON. EXP-INTERNATIONAL-18%", "ledger", "B"),
    ("TRAVELLING AND CONVEYANCE (DOLLAR)", "ledger", "B"),
    ("TRAVELLING AND CONVEYANCE EXPENSE", "ledger", "B"),
    ("TRAVELLING AND CONVEYANCE EXPENSE- 18%", "ledger", "B"),
    ("TRAVELLING AND CONVEYANCE EXPENSE- 18%-DIRECTOR", "ledger", "B"),
    ("TRAVELLING AND CONVEYANCE EXPENSE- 18% (SERVICE)", "ledger", "B"),
    ("TRAVELLING AND CONVEYANCE EXPENSE-DIRECTOR", "ledger", "B"),
    ("TRAVELLING EXP - 5%", "ledger", "B"),
    ("TRAVELLING EXP - 5%-DIRECTOR", "ledger", "B"),
    ("VEHICAL EXP", "ledger", "B"),
    ("CAR & VEHICLE EXP ACCOUNTS (NON-GST)", "ledger", "B"),
    ("PETROL & DIESEL EXP ACCOUNTS", "ledger", "B"),
    ("ENVIRONMENTAL PROTECTION", "ledger", "B"),
    ("ENVIRONMENTAL PROTECTION EXPENSES A/C", "ledger", "B"),
    ("ENVIRONMENTAL PROTECTION EXPENSES A/C-5%", "ledger", "B"),
    ("FINANCE COST", "ledger", "B"),
    ("BANK CHARGES-IMPORT DIRECT PAYMENT", "ledger", "B"),
    ("BANK CHARGES ACC-IMPORT DIRECT PAYMENT", "ledger", "B"),
    ("COMMISSION CHARGES-IMPORT PAYMENT", "ledger", "B"),
    ("CURRENCY CONVERSION TAX CGST", "ledger", "B"),
    ("CURRENCY CONVERSION TAX SGST", "ledger", "B"),
    ("GFID CHARGES-IMPORT", "ledger", "B"),
    ("LOAN PROCESSING CHARGES ACCOUNTS-ITC", "ledger", "B"),
    ("SWIFT CHARGES-IMPORT", "ledger", "B"),
    ("BANK CHARGES-LC-BC", "ledger", "B"),
    ("BANK CHARGES-LC/BC", "ledger", "B"),
    ("COMMISSION CHARGES-LC-BC PAYMENT", "ledger", "B"),
    ("COMMSSION IN LIEU OF EXCHANGE-LC-BC", "ledger", "B"),
    ("INTEREST EXP ON LC/BC PAYMENT", "ledger", "B"),
    ("SWIFT CHARGES-LC-BC", "ledger", "B"),
    ("INTEREST COST", "ledger", "B"),
    ("BANK INTEREST ACCOUNTS", "ledger", "B"),
    ("INTEREST ON ESIC/PF", "ledger", "B"),
    ("INTEREST ON GST PAYMENT", "ledger", "B"),
    ("INTEREST ON LOAN (CAR LOAN)", "ledger", "B"),
    ("INTEREST ON LOAN -TERM LOAN", "ledger", "B"),
    ("INTEREST ON TDS-TCS", "ledger", "B"),
    ("INTEREST ON  UNSECURED LOAN A/C", "ledger", "B"),
    ("OTHER BANK CHARGES", "ledger", "B"),
    ("BANK CHARGES ACCOUNTS", "ledger", "B"),
    ("CREDIT CARD CHARGES", "ledger", "B"),
    ("LOAN PROCESSING CHARGES ACCOUNTS", "ledger", "B"),
    ("INSURANCE EXP", "ledger", "B"),
    ("CAR & VEHICLE INSURANCE EXP ACCOUNTS", "ledger", "B"),
    ("CAR & VEHICLE INSURANCE EXP ACCOUNTS@18%", "ledger", "B"),
    ("INSURANCE EXPENSE ACCOUNT@18%", "ledger", "B"),
    ("INSURANCE EXP-NON GST", "ledger", "B"),
    ("MARINE INSURANCE EXP - 18%", "ledger", "B"),
    ("MARINE INSURANCE EXP - 18%- (PAPER)", "ledger", "B"),
    ("MARINE INSURANCE EXP-NON GST", "ledger", "B"),
    ("TRAVELLING INSURANCE", "ledger", "B"),
    ("LABORATORY EXPENCES", "ledger", "B"),
    ("LABORATORY EXPENSE", "ledger", "B"),
    ("LABORATORY EXPENSE-18%", "ledger", "B"),
    ("LABORATORY EXPENSE-5%", "ledger", "B"),
    ("Learning and Development", "ledger", "B"),
    ("TRAINING EXPENSES-18%", "ledger", "B"),
    ("LEGAL AND PROFESSIONAL EXP", "ledger", "B"),
    ("CONSULATANCY CHARGE", "ledger", "B"),
    ("CONSULTANCY CHARGES-NON GST", "ledger", "B"),
    ("CONSULTANCY CHARGES-RCM", "ledger", "B"),
    ("LEGAL AND PROFESSIONAL A/C(ITC)", "ledger", "B"),
    ("LEGAL FEES(RCM)", "ledger", "B"),
    ("PROFESSIONAL FEES A/C@18%", "ledger", "B"),
    ("TECHNICAL & CONSULTANCY FEES (ITC)", "ledger", "B"),
    ("TRADEMARK COUNSULTANCY (ITC)", "ledger", "B"),
    ("OFFICE EXPENSES", "ledger", "B"),
    ("OTHER EXPENSES", "ledger", "B"),
    ("ROC TAXES", "ledger", "B"),
    ("OFFICE CONSUMABLES EXP", "ledger", "B"),
    ("OFFICE EXP ACCOUNTS", "ledger", "B"),
    ("OFFICE EXPENSES 18%", "ledger", "B"),
    ("OFFICE EXPENSES 5%", "ledger", "B"),
    ("OFFICE EXPENSES @12%", "ledger", "B"),
    ("OFFICE MAINTANCE EXP ACCOUNT", "ledger", "B"),
    ("OFFICE TITAANIUM AREA REPARING", "ledger", "B"),
    ("ROUND OFF", "ledger", "B"),
    ("WATER CHARGES", "ledger", "B"),
    ("PRINTING AND COURIER", "ledger", "B"),
    ("COURIER & POSTAGE EXP", "ledger", "B"),
    ("POSTAGE AND COURIER EXPENSES - 18%", "ledger", "B"),
    ("POSTAGE AND COURIER EXPENSES -HEAD", "ledger", "B"),
    ("POSTAGE AND COURIER EXPENSES-(NON GST)", "ledger", "B"),
    ("POSTAGE AND COURIER EXPENSES -SPARE", "ledger", "B"),
    ("POSTAGE & COURIER EXP (NON-GST)", "ledger", "B"),
    ("PRINTING STATIONERY EXP", "ledger", "B"),
    ("PRINTING & STATIONERY @12%", "ledger", "B"),
    ("PRINTING & STATIONERY @18%", "ledger", "B"),
    ("PRINTING & STATIONERY EXP ACCOUNTS", "ledger", "B"),
    ("RENT EXPENSES", "ledger", "B"),
    ("GODOWN RENT", "ledger", "B"),
    ("GODOWN RENT-RCM", "ledger", "B"),
    ("RENT-TITAANIUM", "ledger", "B"),
    ("REPAIR & MAINTENANCE COST", "ledger", "B"),
    ("ELECTRICAL EXPENSES", "ledger", "B"),
    ("ELECTRICAL EXPENSES  - 18%", "ledger", "B"),
    ("ELECTRICAL EXPENSES ACCOUNTS", "ledger", "B"),
    ("REPAIR & MAINTENANCE EXP (OFFICE)", "ledger", "B"),
    ("FIT OUT CHARGES (MUMBAI OFFICE)", "ledger", "B"),
    ("HOJIWALA GODOWN  REPAIRS", "ledger", "B"),
    ("HOJIWALA GODOWN REPAIRS-12%", "ledger", "B"),
    ("HOJIWALA GODOWN REPAIRS-18%", "ledger", "B"),
    ("HVAC CHARGES (MUMBAI OFFICE)", "ledger", "B"),
    ("MAINTENANCE EXP-TITAANIUM", "ledger", "B"),
    ("MUMBAI OFFICE MAINTENANCE ACCOUNTS", "ledger", "B"),
    ("OFFICE TITAANIUM AREA REPARING-18%", "ledger", "B"),
    ("REPAIR & MAINTENANCE CHARGES", "ledger", "B"),
    ("REPAIR & MAINTENANCE CHARGES@18%", "ledger", "B"),
    ("REPAIR & MAINTENANCE CHARGES@28%", "ledger", "B"),
    ("SACHIN GODOWN REPAIRS-18%", "ledger", "B"),
    ("SECURITY EXP", "ledger", "B"),
    ("SECURITY EXP(RCM)", "ledger", "B"),
    ("SUBSCRIPTION & MEMBERSHIP EXP", "ledger", "B"),
    ("AI SUBSCRIPTION", "ledger", "B"),
    ("AI SUBSCRIPTION EXPENSE (RCM)", "ledger", "B"),
    ("BUSINESS PROMOTION EXPENSE-18%", "ledger", "B"),
    ("REGISTRATION CHARGES", "ledger", "B"),
    ("SUBSCRIPTION EXPENSE (RCM)", "ledger", "B"),
    ("SUBSCRIPTION & MEMBERSHIP EXP (ITC)", "ledger", "B"),
    ("SUBSCRIPTION & MEMBERSHIP EXP (NON-GST)", "ledger", "B"),
    ("TAXES EXP", "ledger", "B"),
    ("GST EXPENSE", "ledger", "B"),
    ("GST LATE FEES", "ledger", "B"),
    ("GUJARAT POLLUTION CONTROL BOARD", "ledger", "B"),
    ("PROPERTY TAX", "ledger", "B"),
    ("STAMP PURCHASE AND NOTERY EXP ACCOUNTS", "ledger", "B"),
    ("TELEPHONE EXP", "ledger", "B"),
    ("TELEPHONE & MOBILE EXPENSES - 18%", "ledger", "B"),
    ("TELEPHONE & MOBILE EXPENSES(NON-GST)", "ledger", "B"),
    ("TRANSPORTATION EXP", "ledger", "B"),
    ("CARRIAGE OUTWARD ACCOUNTS", "ledger", "B"),
    ("CARRIAGE OUTWARD ACCOUNTS - 12%", "ledger", "B"),
    ("CARRIAGE OUTWARD ACCOUNTS (RCM)", "ledger", "B"),
    ("CARRIAGE OUTWARD EXP ACCOUNTS@5%", "ledger", "B"),
    ("CARRIRAGE OUTWARD ACCOUNTS- 18%", "ledger", "B"),
    ("LOADING & UNLOADING EXP", "ledger", "B"),
    ("LOADING & UNLOADING ACCOUNTS@18%", "ledger", "B"),
    ("LOADING & UNLOADING CHARGES (RCM)", "ledger", "B"),
    ("LOADING & UNLOADING EXP-(NON RCM)", "ledger", "B"),
    ("LOADING & UNLOADING EXP (PAPER)", "ledger", "B"),
    ("Excess Short Income Tax 24-25", "ledger", "B"),
    ("Nett Profit:", "group_c", "C"),
]

# ---------------------------------------------------------------------------
# Balance Sheet line-item order (from Files/PL FOMRAT.xlsx, "BS " sheet, rows 5-55)
# ---------------------------------------------------------------------------
BS_ACCOUNT_ORDER = [
    ("Sources of Funds:", "header", "C"),
    ("Capital Account", "group_c", "C"),
    ("Reserves & Surplus", "ledger", "B"),
    ("PREMIUM A/C", "ledger", "B"),
    ("RESERVE & SURPLUS", "ledger", "B"),
    ("SHARE CAPITAL A/C", "ledger", "B"),
    ("Loans (Liability)", "group_c", "C"),
    ("Bank OD A/c", "ledger", "B"),
    ("Secured Loans", "ledger", "B"),
    ("Unsecured Loans", "ledger", "B"),
    ("BANK LC-BC LIMIT", "ledger", "B"),
    ("Current Liabilities", "group_c", "C"),
    ("Duties & Taxes", "ledger", "B"),
    ("Provisions", "ledger", "B"),
    ("Sundry Creditors", "ledger", "B"),
    ("Balance with Related Party", "group_b", "B"),
    ("UNIVERSAL IMPEX", "ledger", "B"),
    ("COLORIX DIGITAL PRINTING SOLUTIONS LLP", "ledger", "B"),
    ("ADVANCE FOR CAPITAL GOODS", "ledger", "B"),
    ("Other Current Liabilities", "ledger", "B"),
    ("AUDIT &PROFESSIONAL FEE PAYABLE", "ledger", "B"),
    ("Branch / Divisions", "group_c", "C"),
    ("Suspense A/c", "group_c", "C"),
    ("PROFIT & LOSS A/C", "group_c", "C"),
    ("Opening Balance", "ledger", "B"),
    ("Current Period", "ledger", "B"),
    ("Unadjusted Forex Gain/Loss", "ledger", "B"),
    ("Total", "subtotal_sources", "C"),
    ("Application of Funds:", "header", "C"),
    ("Fixed Assets", "group_c", "C"),
    ("Investments", "group_c", "C"),
    ("Current Assets", "group_c", "C"),
    ("Closing Stock", "ledger", "B"),
    ("Deposits (Asset)", "ledger", "B"),
    ("Loans & Advances (Asset)", "ledger", "B"),
    ("Balance with Related Party Advances", "ledger", "B"),
    ("Sundry Debtors", "ledger", "B"),
    ("Balance with Related Party Deb", "group_b", "B"),
    ("COLORIX DIGITAL PRINTING SOLUTIONS LLP-SALES", "ledger", "B"),
    ("UNIVERSAL DYE CHEM PVT LTD", "ledger", "B"),
    ("Cash-in-hand", "ledger", "B"),
    ("Bank Accounts", "ledger", "B"),
    ("Advance Tax", "ledger", "B"),
    ("Contingent Assets", "ledger", "B"),
    ("Inventory Valn & Adj", "ledger", "B"),
    ("TDS Receivables", "ledger", "B"),
    ("TDS TCS Receivable", "ledger", "B"),
    ("Difference in opening balances", "ledger", "B"),
    ("INTEREST RECEIVABLE", "ledger", "B"),
    ("PREPAID COMPUTER & PRINTER EXPENSE", "ledger", "B"),
    ("Total", "subtotal_application", "C"),
]

# Alias fallback for group names Tally may store under a different exact label
GROUP_NAME_ALIASES = {
    "PROFIT & LOSS A/C": ["PROFIT & LOSS ACCOUNT"],
}

# Known header info for companies present in Files/PL FOMRAT.xlsx.
# Companies not listed here fall back to name + period only.
COMPANY_HEADER_INFO = {
    "ORANGE O TEC PRIVATE LIMITED": (
        "7111/A2,ROAD NO.71, G.I.D.C,SACHIN SURAT - 394230 | "
        "Delivery: PLOT NO.C-15/3, HOJIWALA INDUSTRIAL ESTATE, ROAD NO.14, HOJIWALA SURAT-394230 | "
        "GST: 24AAHCM0016A1ZD | Email: Accounts@orangeotec.Com | Co.No.85110 94943 | "
        "MSME: UDYAM-GJ-22-0028640-MEDIUM"
    ),
    "ORANGE O TEC PRIVATE LIMITED-NOIDA": (
        "PLOT NO-B-115,SECTOR-67, GAUTAM BUDH NAGAR, NOIDA - 201301 | "
        "GST NO: 09AAHCM0016A1Z5 | Email: delhioffice@orangeotec.com | "
        "MSME-UDYAM-GJ-22-0028640-MEDIUM"
    ),
    "ORANGE O TEC ENTERPRISES PVT LTD": (
        "7111/A2, ROAD NO 71,GATE NO.01, G.I.D.C SACHIN, SURAT-394230(GUJ) INDIA | "
        "Email: Accounts@orangeotec.Com | MSME: UDYAM-GJ-22-0366110-MICRO"
    ),
    "ORANGE O TEC ENTERPRISES PRIVATE LIMITED-NOIDA": (
        "PLOT NO.97,ECCOTECH-12, GRETER NOIDA,GAUTAM BUDH NAGAR,-201308 | "
        "COMPANY GSTIN: 09AAECO3423C1ZT"
    ),
}


def header_info_for(company: str) -> str:
    for key, info in COMPANY_HEADER_INFO.items():
        if key in company:
            return info
    return ""


def post_to_tally(xml_body: str, host: str = TALLY_HOST, port: int = TALLY_PORT) -> str:
    url = f"http://{host}:{port}"
    encoded = xml_body.encode("utf-8")
    req = urllib.request.Request(url, data=encoded, headers={"Content-Type": "text/xml; charset=utf-8"})
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            return resp.read().decode("utf-8", errors="replace")
    except urllib.error.URLError as e:
        raise RuntimeError(
            f"Cannot connect to Tally at {url}.\n"
            f"  Cause: {e}\n"
            f"  Make sure Tally is running, the HTTP XML server is enabled,\n"
            f"  and at least one company is open."
        ) from e


def get_open_companies(host: str = TALLY_HOST, port: int = TALLY_PORT) -> list:
    xml = """<ENVELOPE>
  <HEADER>
    <VERSION>1</VERSION>
    <TALLYREQUEST>Export</TALLYREQUEST>
    <TYPE>Collection</TYPE>
    <ID>AllCompanies</ID>
  </HEADER>
  <BODY>
    <DESC>
      <STATICVARIABLES>
        <SVEXPORTFORMAT>$$SysName:XML</SVEXPORTFORMAT>
      </STATICVARIABLES>
      <TDL>
        <TDLMESSAGE>
          <COLLECTION ISMODIFY="No" NAME="AllCompanies">
            <TYPE>Company</TYPE>
            <FETCH>Name,StartingFrom,EndingAt</FETCH>
          </COLLECTION>
        </TDLMESSAGE>
      </TDL>
    </DESC>
  </BODY>
</ENVELOPE>"""
    resp = post_to_tally(xml, host, port)
    companies = []
    try:
        root = ET.fromstring(resp)
        for elem in root.iter("COMPANY"):
            name_elem = elem.find("NAME")
            if name_elem is not None and name_elem.text and name_elem.text.strip():
                name = name_elem.text.strip()
                if name not in companies:
                    companies.append(name)
    except ET.ParseError:
        pass
    if not companies:
        print("WARNING: Could not detect open companies from Tally.", file=sys.stderr)
    return companies


def parse_tally_amount(element) -> float:
    if element is None:
        return 0.0
    text = (element.text or "").strip()
    if not text:
        return 0.0
    try:
        amount = float(text)
    except ValueError:
        return 0.0
    type_attr = (element.get("TYPE") or "").upper()
    return amount if type_attr != "DR" else -amount


def fetch_balances(from_date: date, to_date: date, company: str, pl_only: bool,
                    host: str = TALLY_HOST, port: int = TALLY_PORT):
    """Fetch ledger + group closing balances for a company/period.
    pl_only=True -> P&L-affecting items only (for the P&L sheet).
    pl_only=False -> non-P&L-affecting items only (for the Balance Sheet).
    """
    from_str = from_date.strftime("%Y%m%d")
    to_str = to_date.strftime("%Y%m%d")
    filter_tag = "IsPLAffecting" if pl_only else "NOT IsPLAffecting"

    xml = f"""<ENVELOPE>
  <HEADER><TALLYREQUEST>Export Data</TALLYREQUEST></HEADER>
  <BODY>
    <EXPORTDATA>
      <REQUESTDESC>
        <STATICVARIABLES>
          <SVFROMDATE>{from_str}</SVFROMDATE>
          <SVTODATE>{to_str}</SVTODATE>
          <SVEXPORTFORMAT>$$SysName:XML</SVEXPORTFORMAT>
          <SVCURRENTCOMPANY>{company}</SVCURRENTCOMPANY>
        </STATICVARIABLES>
        <FORMULACOLLECTION>
          <COLLECTION NAME="LEDGERS_C">
            <TYPE>Ledger</TYPE>
            <FETCH>Name,Parent,ClosingBalance</FETCH>
            <FILTER>{filter_tag}</FILTER>
          </COLLECTION>
          <COLLECTION NAME="GROUPS_C">
            <TYPE>Group</TYPE>
            <FETCH>Name,ClosingBalance</FETCH>
            <FILTER>{filter_tag}</FILTER>
          </COLLECTION>
        </FORMULACOLLECTION>
      </REQUESTDESC>
    </EXPORTDATA>
  </BODY>
</ENVELOPE>"""
    resp = post_to_tally(xml, host, port)
    ledger_data, group_data = {}, {}
    try:
        root = ET.fromstring(resp)
        for ledger in root.iter("LEDGER"):
            name_el = ledger.find("NAME")
            bal_el = ledger.find("CLOSINGBALANCE")
            if name_el is not None and name_el.text:
                ledger_data[name_el.text.strip().upper()] = parse_tally_amount(bal_el)
        for grp in root.iter("GROUP"):
            name_el = grp.find("NAME")
            bal_el = grp.find("CLOSINGBALANCE")
            if name_el is not None and name_el.text:
                group_data[name_el.text.strip().upper()] = parse_tally_amount(bal_el)
    except ET.ParseError as e:
        print(f"  XML parse error for {company}: {e}", file=sys.stderr)
    return ledger_data, group_data


def lookup(name: str, ledger_data: dict, group_data: dict) -> float:
    key = name.strip().upper()
    if key in ledger_data:
        return ledger_data[key]
    if key in group_data:
        return group_data[key]
    for alias in GROUP_NAME_ALIASES.get(name.strip(), []):
        akey = alias.upper()
        if akey in ledger_data:
            return ledger_data[akey]
        if akey in group_data:
            return group_data[akey]
    for prefix in ("ADD: ", "LESS: ", "ADD:", "LESS:"):
        stripped = key.removeprefix(prefix).strip()
        if stripped in ledger_data:
            return ledger_data[stripped]
        if stripped in group_data:
            return group_data[stripped]
    return 0.0


def compute_pl_amounts(ledger_data: dict, group_data: dict) -> list:
    n = len(PL_ACCOUNT_ORDER)
    amounts = [0.0] * n
    idx = {name: i for i, (name, _, _) in enumerate(PL_ACCOUNT_ORDER)}

    for i, (name, rtype, _) in enumerate(PL_ACCOUNT_ORDER):
        if rtype in ("ledger", "group_c", "group_b"):
            amounts[i] = lookup(name, ledger_data, group_data)

    def find(label, occurrence=0):
        hits = [i for i, (nm, _, _) in enumerate(PL_ACCOUNT_ORDER) if nm == label]
        return hits[occurrence] if len(hits) > occurrence else None

    i_sales = find("Sales Accounts")
    i_direct_inc = find("Direct Incomes")
    i_gross_rev_subtotal = find("", 0)  # first blank subtotal_c after Direct Incomes block
    i_cost_of_sales = find("Cost of Sales :")
    i_opening = find("Opening Stock")
    i_purchase = find("Add: Purchase Accounts")
    i_closing = find("Less: Closing Stock")
    i_cos_subtotal = find("", 1)  # second blank -> subtotal_b
    i_direct_exp = find("Direct Expenses")
    i_gross_profit = find("Gross Profit :")
    i_indirect_inc = find("Indirect Incomes")
    i_indirect_inc_subtotal = find("", 2)
    i_indirect_exp = find("Indirect Expenses")
    i_net_profit = find("Nett Profit:")
    i_excess_tax = find("Excess Short Income Tax 24-25")

    if i_gross_rev_subtotal is not None:
        amounts[i_gross_rev_subtotal] = amounts[i_sales] + amounts[i_direct_inc]

    if i_cos_subtotal is not None:
        amounts[i_cos_subtotal] = amounts[i_opening] + amounts[i_purchase] - amounts[i_closing]

    if i_cost_of_sales is not None and amounts[i_cost_of_sales] == 0.0:
        amounts[i_cost_of_sales] = amounts[i_cos_subtotal] + abs(amounts[i_direct_exp])

    if i_gross_profit is not None and amounts[i_gross_profit] == 0.0:
        amounts[i_gross_profit] = amounts[i_gross_rev_subtotal] - amounts[i_cost_of_sales]

    if i_indirect_inc_subtotal is not None:
        amounts[i_indirect_inc_subtotal] = amounts[i_gross_profit] + amounts[i_indirect_inc]

    if i_net_profit is not None and amounts[i_net_profit] == 0.0:
        excess_tax = amounts[i_excess_tax] if i_excess_tax is not None else 0.0
        amounts[i_net_profit] = amounts[i_indirect_inc_subtotal] - abs(amounts[i_indirect_exp]) - excess_tax

    return amounts


def compute_bs_amounts(ledger_data: dict, group_data: dict) -> list:
    n = len(BS_ACCOUNT_ORDER)
    amounts = [0.0] * n

    for i, (name, rtype, _) in enumerate(BS_ACCOUNT_ORDER):
        if rtype in ("ledger", "group_c", "group_b"):
            amounts[i] = lookup(name, ledger_data, group_data)

    def find(label):
        for i, (nm, _, _) in enumerate(BS_ACCOUNT_ORDER):
            if nm == label:
                return i
        return None

    sources_components = ["Capital Account", "Loans (Liability)", "Current Liabilities",
                           "Balance with Related Party", "Branch / Divisions", "Suspense A/c",
                           "PROFIT & LOSS A/C"]
    application_components = ["Fixed Assets", "Investments", "Current Assets"]

    i_total_sources = find("Total")
    # Second "Total" is the application-side total; find both occurrences explicitly.
    total_positions = [i for i, (nm, rtype, _) in enumerate(BS_ACCOUNT_ORDER)
                        if nm == "Total"]
    i_total_sources = total_positions[0] if len(total_positions) > 0 else None
    i_total_application = total_positions[1] if len(total_positions) > 1 else None

    if i_total_sources is not None:
        amounts[i_total_sources] = sum(lookup(c, ledger_data, group_data) for c in sources_components)
    if i_total_application is not None:
        amounts[i_total_application] = sum(lookup(c, ledger_data, group_data) for c in application_components)

    return amounts


def parse_date(s: str) -> date:
    for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Cannot parse date: {s!r}. Use DD-MM-YYYY or YYYY-MM-DD.")


BOLD_FONT = Font(bold=True, name="Calibri", size=10)
NORMAL_FONT = Font(bold=False, name="Calibri", size=10)
SMALL_ITALIC = Font(italic=True, name="Calibri", size=8)
NUM_FORMAT = "#,##0.00"
COL_A_WIDTH = 55
DATA_COL_WIDTH = 16
BOLD_RTYPES = {"group_c", "group_b", "subtotal_sources", "subtotal_application", "header"}
HEADER_ROWS = 6  # rows 1-6 reserved for title/period/company header block


def write_sheet(ws, title: str, period_label: str, companies: list, account_order: list,
                 company_amounts: dict):
    ws.column_dimensions["A"].width = COL_A_WIDTH
    for c_idx in range(len(companies)):
        ws.column_dimensions[get_column_letter(2 + c_idx)].width = DATA_COL_WIDTH

    ws.cell(1, 1, value=title).font = Font(bold=True, size=12, name="Calibri")
    ws.cell(2, 1, value=f"Period: {period_label}").font = NORMAL_FONT

    ws.cell(4, 1, value="Particulars").font = BOLD_FONT
    for c_idx, company in enumerate(companies):
        col = 2 + c_idx
        name_cell = ws.cell(4, col, value=company)
        name_cell.font = BOLD_FONT
        name_cell.alignment = Alignment(horizontal="center", wrap_text=True)

        info = header_info_for(company)
        info_cell = ws.cell(5, col, value=info if info else "")
        info_cell.font = SMALL_ITALIC
        info_cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="top")
    ws.row_dimensions[5].height = 45

    ws.freeze_panes = f"B{HEADER_ROWS + 2}"

    start_row = HEADER_ROWS + 1
    for i, (name, rtype, _col) in enumerate(account_order):
        row_num = start_row + i
        is_bold = rtype in BOLD_RTYPES
        ws.cell(row_num, 1, value=name if name else None).font = BOLD_FONT if is_bold else NORMAL_FONT
        for c_idx, company in enumerate(companies):
            amount = company_amounts[company][i]
            if abs(amount) < 0.005:
                continue
            cell = ws.cell(row_num, 2 + c_idx, value=amount)
            cell.number_format = NUM_FORMAT
            cell.font = BOLD_FONT if is_bold else NORMAL_FONT


def parse_args():
    p = argparse.ArgumentParser(description="Fetch P&L + BS from Tally for all open companies, side-by-side layout")
    p.add_argument("--from-date", default="01-04-2026", help="Start date DD-MM-YYYY")
    p.add_argument("--to-date", default="30-06-2026", help="End date DD-MM-YYYY")
    p.add_argument("--output", default="Files/PL_BS_Combined_Output.xlsx", help="Output Excel path")
    p.add_argument("--port", type=int, default=TALLY_PORT, help="Tally XML server port")
    return p.parse_args()


def main():
    args = parse_args()
    from_date = parse_date(args.from_date)
    to_date = parse_date(args.to_date)
    period_label = f"{from_date.strftime('%d-%b-%y')} to {to_date.strftime('%d-%b-%y')}"

    print(f"Connecting to Tally on port {args.port}...")
    companies = get_open_companies(host=TALLY_HOST, port=args.port)
    if not companies:
        print("ERROR: No companies detected in Tally. Open at least one company and retry.")
        sys.exit(1)
    print(f"Found {len(companies)} company/companies: {', '.join(companies)}")

    pl_amounts, bs_amounts = {}, {}
    for company in companies:
        print(f"\nProcessing: {company}")
        print("  Fetching P&L balances...", end="", flush=True)
        pl_ledger, pl_group = fetch_balances(from_date, to_date, company, pl_only=True, port=args.port)
        pl_amounts[company] = compute_pl_amounts(pl_ledger, pl_group)
        net_i = next((i for i, (nm, _, _) in enumerate(PL_ACCOUNT_ORDER) if nm == "Nett Profit:"), None)
        print(f" Net Profit: {pl_amounts[company][net_i]:,.2f}" if net_i is not None else "")

        print("  Fetching BS balances...", end="", flush=True)
        bs_ledger, bs_group = fetch_balances(from_date, to_date, company, pl_only=False, port=args.port)
        bs_amounts[company] = compute_bs_amounts(bs_ledger, bs_group)
        tot_positions = [i for i, (nm, _, _) in enumerate(BS_ACCOUNT_ORDER) if nm == "Total"]
        if len(tot_positions) == 2:
            src_tot = bs_amounts[company][tot_positions[0]]
            app_tot = bs_amounts[company][tot_positions[1]]
            print(f" Sources Total: {src_tot:,.2f}  Application Total: {app_tot:,.2f}"
                  f"  Diff: {src_tot - app_tot:,.2f}")

    print(f"\nWriting Excel to: {args.output}")
    wb = openpyxl.Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    ws_pl = wb.create_sheet(title="PL YTD")
    write_sheet(ws_pl, "Profit & Loss A/c", period_label, companies, PL_ACCOUNT_ORDER, pl_amounts)
    ws_bs = wb.create_sheet(title="BS")
    write_sheet(ws_bs, "Balance Sheet", period_label, companies, BS_ACCOUNT_ORDER, bs_amounts)

    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(args.output)
    except PermissionError:
        raise RuntimeError(f"Cannot save to {args.output!r} — file may be open in Excel.\nClose the file and retry.")

    print("\n=== SUMMARY ===")
    for company in companies:
        net_i = next((i for i, (nm, _, _) in enumerate(PL_ACCOUNT_ORDER) if nm == "Nett Profit:"), None)
        print(f"  {company}: Net Profit = {pl_amounts[company][net_i]:,.2f}")
    print(f"\nFile saved: {args.output}")


if __name__ == "__main__":
    main()
