"""
Builds Files/PL_BS_Combined_Output.xlsx from pre-fetched Tally data (P&L, Balance Sheet,
chart-of-accounts primary-group rollups, and stock-summary closing stock) for all 5
companies open in Tally as of 2026-07-15, for period 01-Apr-2026 to 30-Jun-2026.

Data was fetched live via the tally-prime MCP tools (profit-loss, balance-sheet,
chart-of-accounts, stock-summary, query-database) and is hardcoded below because the
raw ledger-level XML approach proved unreliable (hangs / "Unknown Request") in this
Tally installation, while the MCP tool's own request path worked reliably.

Key finding: the Balance Sheet as pulled via ledger/group closing balances does not
include Closing Stock (inventory is valued separately, not via a ledger account), which
is the primary driver of the Sources-vs-Application imbalance seen per company. Adding
Closing Stock closes most/all of the gap; any residual is shown explicitly as a
"Reconciliation Gap" line rather than hidden.
"""

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from pathlib import Path

import sys
sys.path.insert(0, str(Path(__file__).parent))
from fetch_pandl_bs_combined import PL_ACCOUNT_ORDER, header_info_for  # noqa: E402

PERIOD_LABEL = "01-Apr-2026 to 30-Jun-2026"

COMPANIES = [
    "ORANGE O TEC PRIVATE LIMITED (01-04-25TO31-03-27)",
    "ORANGE O TEC PRIVATE LIMITED-NOIDA-(from 1-Apr-25)",
    "ORANGE O TEC ENTERPRISES PVT LTD(F.Y.2026-27)",
    "ORANGE O TEC ENTERPRISES PRIVATE LIMITED-NOIDA -FY 26-27",
    "COLORIX DIGITAL PRINTING SOLUTIONS LLP - (from 1-Apr-20)",
]

# ---------------------------------------------------------------------------
# P&L data per company: GROUP-level (sub-line detail) and PRIMARY-level (roll-up) balances
# ---------------------------------------------------------------------------
PL_DATA = {
    "ORANGE O TEC PRIVATE LIMITED (01-04-25TO31-03-27)": {
        "group": {
            "TELEPHONE EXP": -128001.53, "INSURANCE EXP": -1514906.8,
            "TRANSPORT CHARGES  (PURCHASE)": 0, "BRANCH SALES": 0,
            "LOADING & UNLOADING EXP": -265000, "INTEREST INCOME": 0,
            "AI SUBSCRIPTION": -28770.26, "TRANSPORTATION EXP": 0,
            "SECURITY EXP": -164461.36, "ELECTRICITY EXPENSES": -680765.09,
            "TAXES EXP": -2846003, "OTHER DIRETC EXPENSES": -800602.9,
            "HOTEL EXPENSES": -404764.58, "PRINTING STATIONERY EXP": -74944.78,
            "COURIER & POSTAGE EXP": -799584.36, "EMPLOYEE BENEFITS EXPENSES": -124761,
            "SAMPL PAPER EXP": 0, "SERVICE EXP": -59330.28, "LABORATORY EXPENCES": 0,
            "REPAIR & MAINTENANCE COST": -13307.5, "RELATED PARTY SALES": 6079764,
            "SAMPLE FABRIC EXPENSES": -155014.38, "PACKING EXP": -103778.3,
            "LEGAL AND PROFESSIONAL EXP": -345844.58, "INTEREST COST": -6477903.34,
            "ELECTRICAL EXPENSES": -4968, "CARRIAGE INWARD EXP": -1139020.77,
            "COURIER INCOME": 53370, "VEHICAL EXP": -807897.47,
            "COMMISSION & BROKERAGE EXP": -3827406,
            "DEPRECIATION & AMORTISATION EXPENSE": 0, "FINANCE COST": 0,
            "BANK CHARGES-LC-BC": -216092.18, "OTHER INCOME": 1462000,
            "REPAIR & MAINTENANCE EXP (OFFICE)": -355702,
            "EXCHANGE RATE GAIN/(LOSS)": 968962.37,
            "STAFF WELFARE EXPENSES": -457613.09, "FREIGHT EXP": -9540,
            "AIR FREIGHT": -1791203.03, "EMPLOYEE SALARY EXPENSES": -17100745,
            "CARRIAGE OUTWARD ACCOUNTS": -713870.71,
            "CUSTOM DUTY EXPENSES": -13278146.7,
            "DISCOUNT AND RATE DIFFERENCE (GST)": -12307866.53,
            "MANAGERIAL REMUNERATION": -2850000, "OTHER BANK CHARGES": -1268107.34,
            "DONATION SUBSCRIPTION": -60000, "OFFICE EXPENSES": -430648.33,
            "ADVERTISING PROMOTION EXP": -732256.15,
            "RELATED PARTY PURCHASE": -1214659.45, "OCEAN FREIGHT": -5669918.38,
            "SUBSCRIPTION & MEMBERSHIP EXP": -1108110.5, "OTHER EXPENSES": 0,
            "TRAVELLING EXP": -3462407.1,
            "BANK CHARGES-IMPORT DIRECT PAYMENT": -348808.53,
            "JOB WORK CHARGES": 0, "FREIGHT INCOME": 0,
            "ENVIRONMENTAL PROTECTION": -314599,
            "DIRECT EXPENSES": -16412307.02,
            "DISCOUNT AND RATE DIFFERENCE (NON GST)": -624416,
            "DIRECT RCM EXPENSS": 0, "COMPUTER & PRINTER EXP": -128522.29,
            "CLEARING & FORWERDING EXP": -1317952.39,
            "LEARNING AND DEVELOPMENT": -100000, "TRADING DISCOUNT A/C": 45550,
            "RENT EXPENSES": -1348550,
        },
        "primary": {
            "INDIRECT EXPENSES": -49155321.87, "INDIRECT INCOMES": 300000,
            "DIRECT INCOMES": 1515370, "PURCHASE ACCOUNTS": -319941342.08,
            "SALES ACCOUNTS": 358369145.67, "DIRECT EXPENSES": -39767851.78,
        },
        "stock_opening": 292508691.08, "stock_closing": 284571181.51,
    },
    "ORANGE O TEC PRIVATE LIMITED-NOIDA-(from 1-Apr-25)": {
        "group": {
            "TELEPHONE EXP": 0, "INSURANCE EXP": -4274, "TAXI RENT EXP": 0,
            "TRANSPORT CHARGES  (PURCHASE)": -9000, "LOADING & UNLOADING EXP": -72768,
            "SECURITY EXP": -49248, "ELECTRICITY EXPENSES": -20398,
            "TAXES EXP": 0, "OTHER DIRETC EXPENSES": -11200,
            "HOTEL EXPENSES": -125402.47, "PRINTING STATIONERY EXP": 0,
            "COURIER & POSTAGE EXP": -52516.39, "SERVICE EXP": -351299,
            "REPAIR & MAINTENANCE COST": -11600, "RELATED PARTY SALES": 3677910,
            "SAMPLE FABRIC EXPENSES": -23100, "SAMPLE INK": 0,
            "LEGAL AND PROFESSIONAL EXP": 18352.08, "INTEREST COST": -26,
            "CARRIAGE INWARD EXP": -10121, "VEHICAL EXP": -14697, "FINANCE COST": 0,
            "OTHER INCOME": 226500, "REPAIR & MAINTENANCE EXP (OFFICE)": -27033.46,
            "STAFF WELFARE EXPENSES": -137092.31, "FREIGHT EXP": 0, "JOB WORK": 0,
            "EMPLOYEE SALARY EXPENSES": -2875788,
            "CARRIAGE OUTWARD ACCOUNTS": -38720, "CUSTOM DUTY EXPENSES": 0,
            "DISCOUNT AND RATE DIFFERENCE (GST)": -2160560.36,
            "OTHER BANK CHARGES": -427.79, "OFFICE EXPENSES": -15746.24,
            "ADVERTISING PROMOTION EXP": -3830.35,
            "RELATED PARTY PURCHASE": -1808650, "SUBSCRIPTION & MEMBERSHIP EXP": 0,
            "TRAVELLING EXP": -756505.55, "FREIGHT INCOME": 59830,
            "DIRECT EXPENSES": 0,
            "DISCOUNT AND RATE DIFFERENCE (NON GST)": -195213,
            "COMPUTER & PRINTER EXP": -57355, "CLEARING & FORWERDING EXP": 0,
            "LEARNING AND DEVELOPMENT": 0, "RENT EXPENSES": -628142.92,
        },
        "primary": {
            "INDIRECT EXPENSES": -4962567.4, "INDIRECT INCOMES": 0,
            "DIRECT INCOMES": 286330, "PURCHASE ACCOUNTS": -53729084.58,
            "SALES ACCOUNTS": 66970026.91, "DIRECT EXPENSES": -404720,
        },
        "stock_opening": 32685329.56, "stock_closing": 29258524.91,
    },
    "ORANGE O TEC ENTERPRISES PVT LTD(F.Y.2026-27)": {
        "group": {
            "INSURANCE EXP": -157909.81, "PRINTING AND COURIER": -6096.77,
            "LABOUR COST": -38904, "TRANSPORT CHARGES  (PURCHASE)": -15008.47,
            "LOADING & UNLOADING EXP": -71226, "TRANSPORTATION EXP": -550,
            "ELECTRICITY EXPENSES": -330415.92, "EMPLOYEE COST": -4327869,
            "TAXES EXP": -897300, "HOTEL EXPENSES": -136542.28,
            "PRINTING STATIONERY EXP": -11648.31,
            "CARRIAGE OUTWARD EXP": -163209.49, "EMPLOYEE BENEFITS EXPENSES": -33323,
            "COURIER & POSTAGE EXP": 0,
            "DISCOUNT AND RATE DIFFERENCE": -7286734.92, "SERVICE EXP": 0,
            "LABORATORY EXPENCES": -182740.95,
            "REPAIR & MAINTENANCE COST": -455674.77, "RELATED PARTY SALES": 730500,
            "PACKING EXP": -159312, "LEGAL AND PROFESSIONAL EXP": 36849.47,
            "INTEREST COST": -286284.46, "CARRIAGE INWARD EXP": -275755.22,
            "VEHICAL EXP": -2639, "COMMISSION & BROKERAGE EXP": -1091604,
            "DEPRECIATION & AMORTISATION EXPENSE": 0,
            "EXCHANGE RATE GAIN/(LOSS)": -26710.45, "JOB WORK INCOME": 0,
            "STAFF WELFARE EXPENSES": -297988.22,
            "AIR FREIGHT": -5712.32, "EMPLOYEE SALARY EXPENSES": -5194465,
            "CUSTOM DUTY EXPENSES": -4083424, "MANAGERIAL REMUNERATION": -1200000,
            "OFFICE EXPENSES": -83243, "ADVERTISING PROMOTION EXP": 0,
            "RELATED PARTY PURCHASE": -5324740, "OCEAN FREIGHT": -483699,
            "SUBSCRIPTION & MEMBERSHIP EXP": -43300, "TRAVELLING EXP": -315010,
            "DIRECT EXPENSES": 4729987.27,
            "CLEARING & FORWERDING EXP": -163380.72, "RENT EXPENSES": -539250,
        },
        "primary": {
            "INDIRECT EXPENSES": -16370592.46, "INDIRECT INCOMES": 231247.81,
            "DIRECT INCOMES": 0, "PURCHASE ACCOUNTS": -79571139.12,
            "SALES ACCOUNTS": 138368883.95, "DIRECT EXPENSES": -797405.06,
        },
        "stock_opening": 71324050.22, "stock_closing": 91179012.34,
    },
    "ORANGE O TEC ENTERPRISES PRIVATE LIMITED-NOIDA -FY 26-27": {
        "group": {
            "INSURANCE EXP": -3188.19, "LABOUR COST": -94800,
            "TRANSPORT CHARGES  (PURCHASE)": -5879.34,
            "LOADING & UNLOADING EXP": -1200, "ELECTRICITY EXPENSES": -71543,
            "TAXES EXP": 0, "HOTEL EXPENSES": 0,
            "PRINTING STATIONERY EXP": -2896.75, "CARRIAGE OUTWARD EXP": -38427,
            "COURIER & POSTAGE EXP": -6152, "EMPLOYEE BENEFITS EXPENSES": 0,
            "DISCOUNT AND RATE DIFFERENCE": -259978, "SERVICE EXP": -1900,
            "REPAIR & MAINTENANCE COST": -1300, "RELATED PARTY SALES": 2292810.25,
            "PACKING EXP": 0, "LEGAL AND PROFESSIONAL EXP": -77849.47,
            "INTEREST COST": -396, "CARRIAGE INWARD EXP": -2762,
            "COMMISSION & BROKERAGE EXP": -2585, "SERVICE CHARGES": 0,
            "STAFF WELFARE EXPENSES": 0, "EMPLOYEE SALARY EXPENSES": -2144125,
            "EMPLOYEE TRAVEL COST": -1895, "OFFICE EXPENSES": -26374,
            "ADVERTISING PROMOTION EXP": 0, "RELATED PARTY PURCHASE": -4432934,
            "TRAVELLING EXP": -61144, "JOB WORK CHARGES": 0, "BANK CHARGES": 0,
            "FREIGHT INCOME": 5720, "DIRECT EXPENSES": 0,
            "COMPUTER & PRINTER EXP": -1998, "RENT EXPENSES": -1638000,
        },
        "primary": {
            "INDIRECT EXPENSES": -4172855.65, "INDIRECT INCOMES": 3.96,
            "DIRECT INCOMES": 5720, "PURCHASE ACCOUNTS": -23912575.96,
            "SALES ACCOUNTS": 20882067.91, "DIRECT EXPENSES": -105391.34,
        },
        "stock_opening": 32959492.39, "stock_closing": 37902468.1,
    },
    "COLORIX DIGITAL PRINTING SOLUTIONS LLP - (from 1-Apr-20)": {
        "group": {
            "LEGAL AND PROFESSIONAL EXP": 0, "FINANCE COST": -16485.87,
            "FREIGHT EXP": 0, "OFFICE EXPENSES": -1.69,
            "TRADING DISCOUNT A/C": 0, "RENT EXPENSES": 0,
        },
        "primary": {
            "INDIRECT EXPENSES": -22199.69, "INDIRECT INCOMES": 0,
            "DIRECT INCOMES": 0, "PURCHASE ACCOUNTS": -67260307.76,
            "SALES ACCOUNTS": 82457115.76, "DIRECT EXPENSES": 0,
        },
        "stock_opening": 105648, "stock_closing": 105648,
    },
}

# ---------------------------------------------------------------------------
# Balance Sheet: flat list of Tally's standard primary groups (universal across companies)
# ---------------------------------------------------------------------------
BS_PRIMARY_ORDER = [
    ("Sources of Funds:", "header"),
    ("Capital Account", "src"),
    ("Reserves & Surplus", "src"),
    ("Secured Loans", "src"),
    ("Unsecured Loans", "src"),
    ("Bank OD A/c", "src"),
    ("Duties & Taxes", "src"),
    ("Provisions", "src"),
    ("Sundry Creditors", "src"),
    ("Current Liabilities (other)", "src"),
    ("Loans (Liability) (other)", "src"),
    ("Branch / Divisions", "src"),
    ("Suspense A/c", "src"),
    ("Profit & Loss A/c (current)", "src"),
    ("Total Sources of Funds", "subtotal_src"),
    ("Application of Funds:", "header"),
    ("Fixed Assets", "app"),
    ("Investments", "app"),
    ("Current Assets (other)", "app"),
    ("Bank Accounts", "app"),
    ("Cash-in-hand", "app"),
    ("Deposits (Asset)", "app"),
    ("Loans & Advances (Asset)", "app"),
    ("Sundry Debtors", "app"),
    ("Closing Stock (Inventory)", "app"),
    ("Misc. Expenses (ASSET)", "app"),
    ("Total Application of Funds", "subtotal_app"),
    ("Reconciliation Gap (see note)", "gap"),
]

BS_KEY_MAP = {
    "Capital Account": "CAPITAL ACCOUNT",
    "Reserves & Surplus": "RESERVES & SURPLUS",
    "Secured Loans": "SECURED LOANS",
    "Unsecured Loans": "UNSECURED LOANS",
    "Bank OD A/c": "BANK OD A/C",
    "Duties & Taxes": "DUTIES & TAXES",
    "Provisions": "PROVISIONS",
    "Sundry Creditors": "SUNDRY CREDITORS",
    "Current Liabilities (other)": "CURRENT LIABILITIES",
    "Loans (Liability) (other)": "LOANS (LIABILITY)",
    "Branch / Divisions": "BRANCH / DIVISIONS",
    "Suspense A/c": "SUSPENSE A/C",
    "Fixed Assets": "FIXED ASSETS",
    "Investments": "INVESTMENTS",
    "Current Assets (other)": "CURRENT ASSETS",
    "Bank Accounts": ["BANK ACCOUNTS", "BANK ACCOUNT"],
    "Cash-in-hand": "CASH-IN-HAND",
    "Deposits (Asset)": "DEPOSITS (ASSET)",
    "Loans & Advances (Asset)": "LOANS & ADVANCES (ASSET)",
    "Sundry Debtors": "SUNDRY DEBTORS",
    "Misc. Expenses (ASSET)": "MISC. EXPENSES (ASSET)",
}

BS_DATA = {
    "ORANGE O TEC PRIVATE LIMITED (01-04-25TO31-03-27)": {
        "primary": {
            "STOCK-IN-HAND": 0, "LOANS & ADVANCES (ASSET)": -91737258.63,
            "BRANCH / DIVISIONS": -80342513.34, "CASH-IN-HAND": -2005567,
            "FIXED ASSETS": -206600903.87, "PROVISIONS": 30408410,
            "SUNDRY DEBTORS": -418438562.68, "UNSECURED LOANS": -10000000,
            "BANK OD A/C": 369652919.89, "": 34056763.5,
            "SUNDRY CREDITORS": 293085268.63, "INVESTMENTS": -92897847.34,
            "CAPITAL ACCOUNT": 402865682.77, "SECURED LOANS": 107881148.27,
            "MISC. EXPENSES (ASSET)": 0, "CURRENT ASSETS": -54339781.18,
            "LOANS (LIABILITY)": 17770460.85, "DEPOSITS (ASSET)": -3114959,
            "CURRENT LIABILITIES": 8050951.31, "BANK ACCOUNTS": -995937.15,
            "DUTIES & TAXES": -13247094.68, "RESERVES & SURPLUS": 30000000,
        },
        "grand_total_gap": 320051180.35,
    },
    "ORANGE O TEC PRIVATE LIMITED-NOIDA-(from 1-Apr-25)": {
        "primary": {
            "LOANS & ADVANCES (ASSET)": -535000, "BRANCH / DIVISIONS": 81492513.34,
            "CASH-IN-HAND": -221879.93, "FIXED ASSETS": -3219687.87,
            "PROVISIONS": 0, "SUNDRY DEBTORS": -56311309.97, "SUSPENSE A/C": 0,
            "": 10257063.86, "SUNDRY CREDITORS": -1895267.4,
            "CURRENT ASSETS": -73823, "DEPOSITS (ASSET)": -429744,
            "CURRENT LIABILITIES": 853318, "BANK ACCOUNTS": -857482.2,
            "DUTIES & TAXES": 199824.08,
        },
        "grand_total_gap": 29258524.91,
    },
    "ORANGE O TEC ENTERPRISES PVT LTD(F.Y.2026-27)": {
        "primary": {
            "LOANS & ADVANCES (ASSET)": -159000, "BRANCH / DIVISIONS": -89794448.8,
            "CASH-IN-HAND": -543776, "FIXED ASSETS": -26393355.06,
            "PROVISIONS": 8667001.01, "SUNDRY DEBTORS": -152949772.8,
            "UNSECURED LOANS": 81862326, "BANK OD A/C": 3917866.21,
            "SUSPENSE A/C": 13132, "": 62394572.24,
            "SUNDRY CREDITORS": 153356873.06, "INVESTMENTS": -13774660,
            "CAPITAL ACCOUNT": 72737069.56, "SECURED LOANS": 11608048.01,
            "CURRENT ASSETS": -19216477.05, "DEPOSITS (ASSET)": -642381,
            "CURRENT LIABILITIES": 2016288, "BANK ACCOUNTS": -4941525,
            "DUTIES & TAXES": 3699846.96,
        },
        "grand_total_gap": 91857627.34,
    },
    "ORANGE O TEC ENTERPRISES PRIVATE LIMITED-NOIDA -FY 26-27": {
        "primary": {
            "LOANS & ADVANCES (ASSET)": -32712, "BRANCH / DIVISIONS": 89794448.8,
            "CASH-IN-HAND": -285605.87, "FIXED ASSETS": -850877.01,
            "PROVISIONS": 5277, "SUNDRY DEBTORS": -31507208, "SUSPENSE A/C": 12750,
            "": -2360055.37, "SUNDRY CREDITORS": 254420, "CAPITAL ACCOUNT": 0,
            "CURRENT ASSETS": -5298311.61, "DEPOSITS (ASSET)": -1560000,
            "CURRENT LIABILITIES": 705147, "BANK ACCOUNTS": -2152315.57,
            "DUTIES & TAXES": -8822489.27,
        },
        "grand_total_gap": 37902468.1,
    },
    "COLORIX DIGITAL PRINTING SOLUTIONS LLP - (from 1-Apr-20)": {
        "primary": {
            "LOANS & ADVANCES (ASSET)": -95547462, "BRANCH / DIVISIONS": 0,
            "CASH-IN-HAND": -500, "PROVISIONS": 0, "SUNDRY DEBTORS": 38755051.43,
            "BANK OD A/C": 0, "SUSPENSE A/C": 0, "BANK ACCOUNT": -81287.1,
            "": 14657595.33, "SUNDRY CREDITORS": 45987792, "INVESTMENTS": -3001131,
            "CAPITAL ACCOUNT": 50000, "CURRENT ASSETS": -129964,
            "CURRENT LIABILITIES": 0, "DUTIES & TAXES": -506618.28,
            "RESERVES & SURPLUS": -123428.38,
        },
        "grand_total_gap": 60048,
    },
}


def pl_lookup(name, rtype, group_data, primary_data):
    key = name.strip().upper()
    if rtype == "group_c":
        if key in primary_data:
            return primary_data[key]
        if key in group_data:
            return group_data[key]
    else:
        if key in group_data:
            return group_data[key]
        if key in primary_data:
            return primary_data[key]
    for prefix in ("ADD: ", "LESS: ", "ADD:", "LESS:"):
        stripped = key.removeprefix(prefix).strip()
        if stripped in group_data:
            return group_data[stripped]
        if stripped in primary_data:
            return primary_data[stripped]
    return 0.0


def compute_pl_amounts(group_data, primary_data, stock_opening, stock_closing):
    n = len(PL_ACCOUNT_ORDER)
    amounts = [0.0] * n

    for i, (name, rtype, _) in enumerate(PL_ACCOUNT_ORDER):
        if rtype in ("ledger", "group_c", "group_b") and name not in ("Opening Stock", "Add: Purchase Accounts", "Less: Closing Stock"):
            amounts[i] = pl_lookup(name, rtype, group_data, primary_data)

    def find(label, occurrence=0):
        hits = [i for i, (nm, _, _) in enumerate(PL_ACCOUNT_ORDER) if nm == label]
        return hits[occurrence] if len(hits) > occurrence else None

    i_sales = find("Sales Accounts")
    i_direct_inc = find("Direct Incomes")
    i_gross_rev_subtotal = find("", 0)
    i_cost_of_sales = find("Cost of Sales :")
    i_opening = find("Opening Stock")
    i_purchase = find("Add: Purchase Accounts")
    i_closing = find("Less: Closing Stock")
    i_cos_subtotal = find("", 1)
    i_direct_exp = find("Direct Expenses")
    i_gross_profit = find("Gross Profit :")
    i_indirect_inc = find("Indirect Incomes")
    i_indirect_inc_subtotal = find("", 2)
    i_indirect_exp = find("Indirect Expenses")
    i_net_profit = find("Nett Profit:")
    i_excess_tax = find("Excess Short Income Tax 24-25")

    amounts[i_opening] = stock_opening
    amounts[i_closing] = stock_closing
    amounts[i_purchase] = abs(pl_lookup("Purchase Accounts", "group_c", group_data, primary_data))
    # "Direct Expenses" is a multi-level Tally group (many sub-groups nested under it);
    # the group-level dict here only sums DIRECT children, undercounting it. Use the
    # true primary-group rollup instead, since this value feeds Cost of Sales / Net Profit.
    amounts[i_direct_exp] = pl_lookup("Direct Expenses", "group_c", group_data, primary_data)

    amounts[i_gross_rev_subtotal] = amounts[i_sales] + amounts[i_direct_inc]
    amounts[i_cos_subtotal] = amounts[i_opening] + amounts[i_purchase] - amounts[i_closing]

    if amounts[i_cost_of_sales] == 0.0:
        amounts[i_cost_of_sales] = amounts[i_cos_subtotal] + abs(amounts[i_direct_exp])

    if amounts[i_gross_profit] == 0.0:
        amounts[i_gross_profit] = amounts[i_gross_rev_subtotal] - amounts[i_cost_of_sales]

    amounts[i_indirect_inc_subtotal] = amounts[i_gross_profit] + amounts[i_indirect_inc]

    if amounts[i_net_profit] == 0.0:
        excess_tax = amounts[i_excess_tax] if i_excess_tax is not None else 0.0
        amounts[i_net_profit] = amounts[i_indirect_inc_subtotal] - abs(amounts[i_indirect_exp]) - excess_tax

    return amounts


def bs_lookup(row_label, primary_data):
    keymap_val = BS_KEY_MAP.get(row_label)
    if keymap_val is None:
        return 0.0
    keys = keymap_val if isinstance(keymap_val, list) else [keymap_val]
    for k in keys:
        if k in primary_data:
            return primary_data[k]
    return 0.0


def compute_bs_amounts(primary_data, stock_closing):
    n = len(BS_PRIMARY_ORDER)
    amounts = [0.0] * n
    for i, (label, kind) in enumerate(BS_PRIMARY_ORDER):
        if kind in ("src", "app") and label != "Closing Stock (Inventory)":
            amounts[i] = bs_lookup(label, primary_data)

    i_closing_stock = next(i for i, (lbl, _) in enumerate(BS_PRIMARY_ORDER) if lbl == "Closing Stock (Inventory)")
    amounts[i_closing_stock] = -stock_closing  # asset => negative per Tally BS sign convention

    i_pl_row = next(i for i, (lbl, _) in enumerate(BS_PRIMARY_ORDER) if lbl == "Profit & Loss A/c (current)")
    amounts[i_pl_row] = primary_data.get("", 0.0)

    i_src_total = next(i for i, (lbl, _) in enumerate(BS_PRIMARY_ORDER) if lbl == "Total Sources of Funds")
    i_app_total = next(i for i, (lbl, _) in enumerate(BS_PRIMARY_ORDER) if lbl == "Total Application of Funds")
    i_gap = next(i for i, (lbl, _) in enumerate(BS_PRIMARY_ORDER) if lbl.startswith("Reconciliation Gap"))

    src_rows = [i for i, (_, kind) in enumerate(BS_PRIMARY_ORDER) if kind == "src"]
    app_rows = [i for i, (_, kind) in enumerate(BS_PRIMARY_ORDER) if kind == "app"]
    amounts[i_src_total] = sum(amounts[i] for i in src_rows)
    amounts[i_app_total] = sum(amounts[i] for i in app_rows)
    amounts[i_gap] = amounts[i_src_total] + amounts[i_app_total]  # should be ~0 if fully reconciled

    return amounts


BOLD_FONT = Font(bold=True, name="Calibri", size=10)
NORMAL_FONT = Font(bold=False, name="Calibri", size=10)
GAP_FONT = Font(bold=True, italic=True, name="Calibri", size=10, color="C00000")
SMALL_ITALIC = Font(italic=True, name="Calibri", size=8)
NUM_FORMAT = "#,##0.00"
COL_A_WIDTH = 55
DATA_COL_WIDTH = 16
HEADER_ROWS = 6


def write_pl_sheet(ws, companies):
    ws.column_dimensions["A"].width = COL_A_WIDTH
    for c_idx in range(len(companies)):
        ws.column_dimensions[get_column_letter(2 + c_idx)].width = DATA_COL_WIDTH

    ws.cell(1, 1, value="Profit & Loss A/c").font = Font(bold=True, size=12, name="Calibri")
    ws.cell(2, 1, value=f"Period: {PERIOD_LABEL}").font = NORMAL_FONT
    ws.cell(4, 1, value="Particulars").font = BOLD_FONT

    company_amounts = {}
    for c_idx, company in enumerate(companies):
        col = 2 + c_idx
        ws.cell(4, col, value=company).font = BOLD_FONT
        ws.cell(4, col).alignment = Alignment(horizontal="center", wrap_text=True)
        info = header_info_for(company)
        ws.cell(5, col, value=info).font = SMALL_ITALIC
        ws.cell(5, col).alignment = Alignment(horizontal="center", wrap_text=True, vertical="top")

        data = PL_DATA[company]
        company_amounts[company] = compute_pl_amounts(
            data["group"], data["primary"], data["stock_opening"], data["stock_closing"]
        )
    ws.row_dimensions[5].height = 45
    ws.freeze_panes = f"B{HEADER_ROWS + 2}"

    BOLD_RTYPES = {"group_c", "group_b", "subtotal_c", "subtotal_b"}
    start_row = HEADER_ROWS + 1
    for i, (name, rtype, _col) in enumerate(PL_ACCOUNT_ORDER):
        row_num = start_row + i
        is_bold = rtype in BOLD_RTYPES
        ws.cell(row_num, 1, value=name if name else None).font = BOLD_FONT if is_bold else NORMAL_FONT
        for c_idx, company in enumerate(companies):
            amount = company_amounts[company][i]
            if abs(amount) < 0.005:
                continue
            cell = ws.cell(row_num, 2 + c_idx, value=amount)
            cell.number_format = NUM_FORMAT
            cell.font = BOLD_FONT if is_bold else NORMAL_FONT

    return company_amounts


def write_bs_sheet(ws, companies):
    ws.column_dimensions["A"].width = COL_A_WIDTH
    for c_idx in range(len(companies)):
        ws.column_dimensions[get_column_letter(2 + c_idx)].width = DATA_COL_WIDTH

    ws.cell(1, 1, value="Balance Sheet (DRAFT - see Reconciliation Gap note)").font = Font(bold=True, size=12, name="Calibri", color="C00000")
    ws.cell(2, 1, value=f"As at: 30-Jun-2026 (period {PERIOD_LABEL})").font = NORMAL_FONT
    ws.cell(4, 1, value="Particulars").font = BOLD_FONT

    company_amounts = {}
    for c_idx, company in enumerate(companies):
        col = 2 + c_idx
        ws.cell(4, col, value=company).font = BOLD_FONT
        ws.cell(4, col).alignment = Alignment(horizontal="center", wrap_text=True)
        info = header_info_for(company)
        ws.cell(5, col, value=info).font = SMALL_ITALIC
        ws.cell(5, col).alignment = Alignment(horizontal="center", wrap_text=True, vertical="top")

        bs = BS_DATA[company]
        company_amounts[company] = compute_bs_amounts(bs["primary"], PL_DATA[company]["stock_closing"])
    ws.row_dimensions[5].height = 45
    ws.freeze_panes = f"B{HEADER_ROWS + 2}"

    start_row = HEADER_ROWS + 1
    for i, (label, kind) in enumerate(BS_PRIMARY_ORDER):
        row_num = start_row + i
        font = BOLD_FONT if kind in ("header", "subtotal_src", "subtotal_app") else (GAP_FONT if kind == "gap" else NORMAL_FONT)
        ws.cell(row_num, 1, value=label).font = font
        for c_idx, company in enumerate(companies):
            amount = company_amounts[company][i]
            if kind == "gap" or abs(amount) >= 0.005:
                cell = ws.cell(row_num, 2 + c_idx, value=amount)
                cell.number_format = NUM_FORMAT
                cell.font = font

    return company_amounts


def main():
    wb = openpyxl.Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    ws_pl = wb.create_sheet(title="PL YTD")
    pl_amounts = write_pl_sheet(ws_pl, COMPANIES)

    ws_bs = wb.create_sheet(title="BS (draft)")
    bs_amounts = write_bs_sheet(ws_bs, COMPANIES)

    # Notes sheet
    ws_notes = wb.create_sheet(title="Notes")
    ws_notes.column_dimensions["A"].width = 100
    notes = [
        "NOTES ON THIS WORKBOOK",
        "",
        "Source: Tally Prime (localhost:9010), all companies open as of 2026-07-15.",
        "Period: 01-Apr-2026 to 30-Jun-2026 (Q1 FY26-27).",
        "",
        "PL YTD sheet: Sales/Direct Income/Purchase/Direct Expense/Indirect Income/Indirect",
        "Expense figures use Tally's own primary-group roll-ups (chart-of-accounts",
        "primary_group), which are authoritative. Sub-line detail below each heading uses",
        "group-level (immediate parent) balances — blank rows mean no data matched that",
        "sub-line for that company, not that the balance is zero elsewhere.",
        "Opening/Closing Stock use the stock-summary tool (inventory valuation), since",
        "these are not simple ledger balances.",
        "",
        "BS (draft) sheet: Balance Sheet ledger/group closing balances, as pulled via the",
        "Tally API, do NOT include inventory valuation (Closing Stock), which is tracked",
        "separately. This causes a Sources-vs-Application imbalance. We added Closing Stock",
        "back in as a Current Asset, which resolves most or all of the gap:",
        "  - ORANGE O TEC PRIVATE LIMITED-NOIDA: gap fully explained by closing stock (exact match)",
        "  - ORANGE O TEC ENTERPRISES PRIVATE LIMITED-NOIDA: gap fully explained by closing stock (exact match)",
        "  - ORANGE O TEC ENTERPRISES PVT LTD: ~99.3% explained, ~Rs 6.8 lakh residual",
        "  - ORANGE O TEC PRIVATE LIMITED: ~88.9% explained, ~Rs 3.5 crore residual",
        "    (this residual figure, Rs 3,54,79,998.84, exactly matches 'Sales Bills to Make'",
        "    from a prior manually-prepared report — likely a goods-in-transit/billing",
        "    adjustment not captured as a standard ledger balance)",
        "  - COLORIX DIGITAL PRINTING SOLUTIONS LLP: small residual (~Rs 45,600), stock did",
        "    not move this quarter so is less likely to be the driver here",
        "",
        "The 'Reconciliation Gap' row on the BS sheet shows this residual explicitly per",
        "company so nothing is silently hidden. Treat the BS sheet as a working draft —",
        "cross-check against Tally's native Balance Sheet report before relying on it for",
        "filings or board reporting.",
    ]
    for i, line in enumerate(notes, start=1):
        cell = ws_notes.cell(i, 1, value=line)
        cell.font = Font(bold=True, size=12) if i == 1 else NORMAL_FONT
        cell.alignment = Alignment(wrap_text=True)

    output_path = Path(__file__).parent.parent.parent.parent.parent / "Files" / "PL_BS_Combined_Output.xlsx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(output_path)
    except PermissionError:
        raise RuntimeError(f"Cannot save to {output_path!r} — file may be open in Excel.\nClose the file and retry.")

    print("=== P&L NET PROFIT SUMMARY ===")
    net_i = next(i for i, (nm, _, _) in enumerate(PL_ACCOUNT_ORDER) if nm == "Nett Profit:")
    for company in COMPANIES:
        print(f"  {company}: {pl_amounts[company][net_i]:,.2f}")

    print("\n=== BS RECONCILIATION GAP SUMMARY ===")
    gap_i = next(i for i, (lbl, _) in enumerate(BS_PRIMARY_ORDER) if lbl.startswith("Reconciliation Gap"))
    src_i = next(i for i, (lbl, _) in enumerate(BS_PRIMARY_ORDER) if lbl == "Total Sources of Funds")
    for company in COMPANIES:
        gap = bs_amounts[company][gap_i]
        src = bs_amounts[company][src_i]
        pct = (abs(gap) / abs(src) * 100) if src else 0
        print(f"  {company}: gap={gap:,.2f} ({pct:.2f}% of sources total {src:,.2f})")

    print(f"\nFile saved: {output_path}")


if __name__ == "__main__":
    main()
