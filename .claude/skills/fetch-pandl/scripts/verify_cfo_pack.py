"""
verify_cfo_pack.py - prove a CFO Pack ties, on values Excel itself recalculates.

    python verify_cfo_pack.py "Files\\Results\\CFO_Pack_PL_BS_CF_FF_<stamp>.xlsx" [--source <PL_BS_GL_Detail.xlsx>]

Why Excel and not the cached values: openpyxl writes formulas with no results, and a stale cache
is exactly how a wrong figure gets trusted. The pack is opened in a private Excel instance on a
COPY, CalculateFull() is run, the copy is saved, and every check below reads THAT copy. (The
`formulas` engine was tried first on the 7 MB source; Excel is the engine the reader will use.)

Checks (each prints PASS / FAIL with the worst offender):
  1. Cash Flow CFS and SFS - closing cash equals the Balance Sheet in every period and column.
  2. Fund Flow CFS and SFS - the working-capital schedule equals Sources less Applications.
  3. Additivity of every Cash Flow line, every column: Apr+May+Jun = Q1, Jul+Aug+Sep = Q2,
     Q1+Q2 = Apr-Sep, Q1+Jul = Apr-Jul, Apr-Jul+Aug = Apr-Aug.
  4. The six copied MOM sheets recalculate to the source workbook's own values.
  5. Opening BS: SFS gap nil for every company.
  6. CFO Analysis: headline figures equal the source P&L / BS read independently; every
     'Checks' row says OK.
  7. The PDF (if present): page count, no #REF!/#VALUE!/#DIV/0!/#NAME? or '####', and the
     Apr-Sep turnover printed on it.
"""
from __future__ import annotations

import argparse
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_cfo_pack as B  # noqa: E402  (layout helpers and constants)

TOL = 1.0  # rupees


def excel_recalc(src: Path) -> Path:
    import pythoncom
    import win32com.client as win32
    tmp = Path(tempfile.mkdtemp(prefix="cfo_verify_")) / src.name
    shutil.copy2(src, tmp)
    pythoncom.CoInitialize()
    xl = win32.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(str(tmp), UpdateLinks=0)
        xl.CalculateFull()
        wb.Save()
        wb.Close(SaveChanges=False)
    finally:
        xl.Quit()
        pythoncom.CoUninitialize()
    return tmp


class Report:
    def __init__(self):
        self.fails = 0

    def check(self, name, ok, detail=""):
        print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f"  - {detail}" if detail else ""))
        if not ok:
            self.fails += 1


def grid(ws):
    return {(r, c): v for r, row in enumerate(ws.iter_rows(values_only=True), start=1)
            for c, v in enumerate(row, start=1) if v is not None}


def find_row(ws_grid, text, col=1):
    for (r, c), v in ws_grid.items():
        if c == col and isinstance(v, str) and v.strip().startswith(text):
            return r
    return None


def header_pos(ws_grid):
    """(period, entity key) -> column, from rows 3 and 4 of a flow sheet."""
    pos, period = {}, None
    cols = sorted(c for (r, c) in ws_grid if r == 4)
    for c in range(2, max(cols) + 1):
        p = ws_grid.get((3, c))
        if p:
            period = p
        e = ws_grid.get((4, c))
        if e and period:
            pos[(period, B.Layout._entity_key(str(e)))] = c
    return pos


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pack")
    ap.add_argument("--source", help="the PL_BS_GL_Detail workbook the pack was built from")
    ap.add_argument("--no-recalc", action="store_true", help="read the pack as saved (already recalculated by Excel)")
    args = ap.parse_args()
    pack = Path(args.pack)
    rep = Report()
    print(f"Verifying {pack.name}")
    calc = pack if args.no_recalc else excel_recalc(pack)
    wb = openpyxl.load_workbook(calc, data_only=True, read_only=True)
    G = {ws.title: grid(ws) for ws in wb.worksheets}

    # 1-3. flow sheets ------------------------------------------------------------------------------
    for sheet, chk_label in (("MOM Cash Flow CFS", "Check: difference"), ("MOM Cash Flow SFS", "Check: difference"),
                             ("MOM Fund Flow CFS", "Check: difference"), ("MOM Fund Flow SFS", "Check: difference")):
        g = G[sheet]
        pos = header_pos(g)
        r = find_row(g, chk_label)
        vals = [(k, g.get((r, c))) for k, c in pos.items()]
        bad = [(k, v) for k, v in vals if not isinstance(v, (int, float)) or abs(v) > TOL]
        rep.check(f"{sheet}: check row nil in all {len(vals)} cells", not bad,
                  f"worst {max(bad, key=lambda x: abs(x[1]) if isinstance(x[1], (int, float)) else 1e18)}" if bad else "")
    for sheet in ("MOM Cash Flow CFS", "MOM Cash Flow SFS"):
        g = G[sheet]
        pos = header_pos(g)
        line_rows = sorted({r for (r, c), v in g.items() if c == 1 and r > 6 and isinstance(v, str)
                            and (r, pos[("Apr-26", "GT")]) in g and not v.startswith(("Check", "Opening cash", "Closing cash"))})
        worst, n = (0.0, None), 0
        rules = [(["Apr-26", "May-26", "Jun-26"], "Q1 Total"), (["Jul-26", "Aug-26", "Sep-26"], "Q2 Total"),
                 (["Q1 Total", "Q2 Total"], "Apr-Sep Total"), (["Q1 Total", "Jul-26"], "Apr-Jul Total"),
                 (["Apr-Jul Total", "Aug-26"], "Apr-Aug Total")]
        for r in line_rows:
            for ek in B.BLOCK8:
                for parts, whole in rules:
                    s = sum((g.get((r, pos[(p, ek)])) or 0.0) for p in parts)
                    w = g.get((r, pos[(whole, ek)])) or 0.0
                    n += 1
                    if abs(s - w) > abs(worst[0]):
                        worst = (s - w, (g.get((r, 1)), ek, whole))
        rep.check(f"{sheet}: additivity of every line ({n} tests)", abs(worst[0]) <= TOL,
                  f"worst {worst[0]:,.2f} at {worst[1]}" if worst[1] else "")

    # 4. copied sheets vs the source ------------------------------------------------------------------
    if args.source:
        src = Path(args.source)
        tmpd = Path(tempfile.mkdtemp(prefix="cfo_src_"))
        sc = tmpd / src.name
        shutil.copy2(src, sc)
        swb = openpyxl.load_workbook(sc, data_only=True, read_only=True)
        for name in B.KEEP_SHEETS:
            a, b = grid(swb[name]), G[name]
            nums = [(k, v) for k, v in a.items() if isinstance(v, (int, float))]
            diff = [(k, v, b.get(k)) for k, v in nums
                    if not isinstance(b.get(k), (int, float)) or abs(b.get(k) - v) > max(TOL, abs(v) * 1e-9)]
            rep.check(f"{name}: {len(nums)} numeric cells equal the source", not diff,
                      f"{len(diff)} differ, e.g. {diff[:3]}" if diff else "")

    # 5. opening ------------------------------------------------------------------------------------
    g = G["Opening BS 31-Mar-26"]
    r = find_row(g, "RECONCILIATION GAP (Liabilities - Assets) - must be 0")
    vals = [g.get((r, c)) or 0.0 for c in range(5, 13)]
    rep.check("Opening BS SFS: gap nil for every company", all(abs(v) <= TOL for v in vals), f"{vals}")
    r2 = [rr for (rr, c), v in g.items() if c == 1 and isinstance(v, str) and v.startswith("RECONCILIATION GAP") and rr != r]
    if r2:
        print(f"  [INFO] Opening BS CFS gap (unmatched inter-company at 31-Mar), Grand Total: {g.get((r2[0], 12)) or 0:,.2f}")

    # 6. CFO analysis -------------------------------------------------------------------------------
    g = G["CFO Analysis"]
    rows_ok = [(r, v) for (r, c), v in g.items() if c == 7 and v in ("OK", "CHECK")]
    bad = [(g.get((r, 1)), g.get((r, 6))) for r, v in rows_ok if v != "OK"]
    rep.check(f"CFO Analysis: all {len(rows_ok)} Checks rows OK", not bad and rows_ok, f"{bad}" if bad else "")
    pl = G["MOM PL GL CFS"]
    lay_cols = header_pos_pl(pl)
    for label_, text in (("TOTAL TURNOVER", "Turnover"), ("EBITDA", "EBITDA"), ("PAT (NET PROFIT)", "PAT")):
        src_r = next(r for (r, c), v in pl.items() if c <= 4 and isinstance(v, str) and v.strip().startswith(label_))
        expect = (pl.get((src_r, lay_cols[("Apr-Sep Total", "GT")])) or 0.0) / B.CR
        rr = find_row(g, text)
        got = g.get((rr, 5))
        rep.check(f"CFO Analysis headline '{text}' Apr-Sep = source ({expect:,.2f} cr)",
                  isinstance(got, (int, float)) and abs(got - expect) < 0.005, f"sheet {got}")
    errs = [(t, k, v) for t, gg in G.items() for k, v in gg.items()
            if isinstance(v, str) and v.startswith("#") and v.rstrip("!?0/").upper() in ("#REF", "#VALUE", "#DIV", "#NAME", "#N/A", "#NUM")]
    rep.check("No Excel error values in any sheet", not errs, f"{len(errs)} e.g. {errs[:5]}" if errs else "")

    # 6b. v4 sheets -------------------------------------------------------------------------------
    for sheet in ("Branch Transactions", "Related Party"):
        g = G.get(sheet)
        if g is None:
            rep.check(f"{sheet} sheet present", False)
            continue
        pos = header_pos_memo(g)
        tots = [r for (r, c), v in g.items() if c == 1 and isinstance(v, str) and v.startswith("TOTAL:")]
        pl_tots = tots[:2]  # the P&L section's sales and purchases totals come first
        net = sum((g.get((r, pos[("Apr-Sep", "GT")])) or 0) for r in pl_tots)
        rep.check(f"{sheet}: P&L sales and purchases net to nil, Grand Total Apr-Sep", abs(net) < 0.01, f"net {net:,.4f} cr")
    # Schedules (debtors / creditors ...) tie to the standalone balance sheet at 30-Sep.
    sch, bs = G.get("Debtors & Creditors"), G["MOM BS GL SFS"]
    if sch is None:
        rep.check("Debtors & Creditors sheet present", False)
    else:
        spos, bpos = header_pos_memo(sch), header_pos_pl(bs)
        for mine, theirs in (("TOTAL: Sundry debtors", "TOTAL: Sundry Debtors"), ("TOTAL: Sundry creditors", "TOTAL: Sundry Creditors")):
            got = sch.get((find_row(sch, mine), spos[("Apr-Sep", "GT")]))
            br = next(r for (r, c), v in bs.items() if c == 4 and v == theirs)
            want = (bs.get((br, bpos[("Sep-26", "GT")])) or 0) / B.CR
            rep.check(f"Debtors & Creditors: {mine[7:]} 30-Sep = balance sheet ({want:,.2f} cr)",
                      isinstance(got, (int, float)) and abs(got - want) < 0.005, f"schedule {got}")
    # Ageing / inventory (as on the snapshot date): the buckets add up, and stock ties to the snapshot.
    if "Inventory Ageing" in G:
        import ageing_pack as A
        snap, _n = A.load_snapshot(B.RESULTS)
        inv, deb = G["Inventory Ageing"], G["Debtors Ageing"]
        tr = find_row(inv, "TOTAL")
        buckets, total = sum(inv.get((tr, c)) or 0 for c in (7, 8, 9)), inv.get((tr, 10)) or 0
        want = sum(s["closing_value"] or 0 for c in snap["companies"].values() for s in c["stock"]) / B.CR
        rep.check(f"Inventory Ageing: buckets = total = snapshot stock ({want:,.2f} cr)",
                  abs(buckets - total) < 0.005 and abs(total - want) < 0.005, f"buckets {buckets:,.4f}, total {total:,.4f}")
        dr = find_row(deb, "TOTAL")
        want = sum(b["pending_amount"] or 0 for c in snap["companies"].values() for b in c["receivables"]) / B.CR
        got = deb.get((dr, 6)) or 0
        rep.check(f"Debtors Ageing: total = snapshot bills ({want:,.2f} cr)", abs(got - want) < 0.005, f"sheet {got:,.4f}")
    m = G["PDF - MOM"]
    blocks, cur = {}, None
    for (r, c), v in sorted(m.items()):
        if c == 1 and isinstance(v, str) and v.startswith("MONTH ON MONTH - ") and "," not in v:
            cur = v[len("MONTH ON MONTH - "):]
        if cur and c == 1 and v == "Turnover":
            blocks[cur] = m.get((r, 10))
    oe, ot, en = (blocks.get(k) for k in ("Orange + Enterprises (CFS)", "Group (CFS)", "Group (CFS)"))
    pl = G["MOM PL GL CFS"]
    cols = header_pos_pl(pl)
    tr = next(r for (r, c), v in pl.items() if c <= 4 and isinstance(v, str) and v.strip().startswith("TOTAL TURNOVER"))
    expect = ((pl.get((tr, cols[("Apr-Sep Total", "TOT_OTEC")])) or 0) + (pl.get((tr, cols[("Apr-Sep Total", "TOT_ENT")])) or 0)) / B.CR
    rep.check(f"Orange + Enterprises view = CFS Total O Tec + Total Enterprises (turnover {expect:,.2f} cr)",
              isinstance(oe, (int, float)) and abs(oe - expect) < 0.005, f"view {oe}")
    dpdf = pack.with_name(pack.name.replace("CFO_Pack_PL_BS_CF_FF_", "Directors_Executive_Summary_")).with_suffix(".pdf")
    if dpdf.exists():
        import pymupdf
        rep.check("Directors' PDF: 3 pages (summary; margin + O Tec / Enterprises; Colorix)", pymupdf.open(dpdf).page_count == 3,
                  f"{pymupdf.open(dpdf).page_count} pages")

    # 7. PDF ------------------------------------------------------------------------------------------
    pdf = pack.with_suffix(".pdf")
    if pdf.exists():
        import pymupdf
        doc = pymupdf.open(pdf)
        text = "\n".join(p.get_text() for p in doc)
        rep.check(f"PDF: {doc.page_count} pages", doc.page_count >= 2)
        bad = [e for e in ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "####") if e in text]
        rep.check("PDF: no error values or #### overflow", not bad, f"{bad}")
        ca = G["CFO Analysis"]
        to = ca.get((find_row(ca, "Turnover"), 5))
        rep.check(f"PDF shows Apr-Sep turnover {to:,.2f}", f"{to:,.2f}" in text)
    print(f"\n{'ALL CHECKS PASSED' if rep.fails == 0 else f'{rep.fails} CHECK(S) FAILED'}")
    print(f"RECALC_COPY={calc}")
    return rep.fails


def header_pos_memo(g):
    """(period, entity key) -> column on the Branch / Related Party sheets (rows 3 and 4)."""
    pos, period = {}, None
    maxc = max(c for (r, c) in g if r == 4)
    for c in range(2, maxc + 1):
        if g.get((3, c)):
            period = str(g.get((3, c))).strip()
        e = g.get((4, c))
        if e and period:
            pos[(period, B.Layout._entity_key(str(e)))] = c
    return pos


def header_pos_pl(g):
    pos, period = {}, None
    maxc = max(c for (r, c) in g if r == 4)
    for c in range(5, maxc + 1):
        p = g.get((3, c))
        if p:
            period = str(p).strip()
        e = g.get((4, c))
        if e and period and "% of Sales" not in str(e):
            pos[(period, B.Layout._entity_key(str(e)))] = c
    return pos


if __name__ == "__main__":
    sys.exit(1 if main() else 0)
