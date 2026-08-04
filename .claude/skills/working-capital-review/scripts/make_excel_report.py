# -*- coding: utf-8 -*-
import argparse
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break

# Derived from this file's location (.../PL and BS/.claude/skills/<skill>/scripts/<file>.py)
# so the script keeps working regardless of which user profile the project sits under.
PROJECT_ROOT = str(Path(__file__).resolve().parents[4])
DEFAULT_BASE = PROJECT_ROOT + r"\Files\CFS 08-07-26.xlsx"
DEFAULT_OUT = PROJECT_ROOT + r"\Files\Working_Capital_Review_Jul2026.xlsx"

parser = argparse.ArgumentParser(description="Build the Working Capital Review workbook from a CFS tracker file.")
parser.add_argument("--base", default=DEFAULT_BASE, help="Path to the source CFS xlsx (default: Files/CFS 08-07-26.xlsx)")
parser.add_argument("--output", default=DEFAULT_OUT, help="Path to write the output workbook")
args = parser.parse_args()

BASE = args.base
OUT = args.output

RED        = "D03B3B"
GREY_GRID  = "E1E0D9"
GREY_MUTED = "898781"
NAVY_TITLE = "1F3864"

# ================= 1. Load base data (values only) — now split across two sheets =================
src_wb = openpyxl.load_workbook(BASE, data_only=True)
src_cur = src_wb[" Jul 26"]   # current period: 8-Jul-2026
src_pri = src_wb["12 June"]   # prior period: 12-Jun-2026

wb = openpyxl.Workbook()
wb.remove(wb.active)

def copy_source_sheet(src_ws, sheet_name, max_row_copy):
    data_ws = wb.create_sheet(sheet_name)
    for row in src_ws.iter_rows(min_row=1, max_row=max_row_copy, min_col=1, max_col=16):
        for cell in row:
            if cell.value is not None:
                data_ws.cell(row=cell.row, column=cell.column, value=cell.value)

    thin_d = Side(style="thin", color="B7B7B7")
    box_d = Border(top=thin_d, bottom=thin_d, left=thin_d, right=thin_d)
    for row in data_ws.iter_rows(min_row=1, max_row=max_row_copy, min_col=2, max_col=16):
        for cell in row:
            if cell.value is not None:
                cell.border = box_d
                if isinstance(cell.value, (int, float)):
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "#,##0"
    for r_ in range(1, max_row_copy + 1):
        b_val = data_ws.cell(row=r_, column=2).value
        if b_val in ("Type", "PRODUCT", "Row Labels") or (
            isinstance(b_val, str) and b_val.strip() in ("Vendor", "Customer", "Inventory")
            and data_ws.cell(row=r_, column=3).value is None
        ):
            for c_ in range(2, 17):
                cell = data_ws.cell(row=r_, column=c_)
                if cell.value is not None or c_ == 2:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor=NAVY_TITLE)
        if isinstance(b_val, str) and "Grand Total" in b_val:
            for c_ in range(2, 17):
                cell = data_ws.cell(row=r_, column=c_)
                if cell.value is not None:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor="D9D9D9")
        elif isinstance(b_val, str) and b_val.strip() in ("Inventory", "Customer", "Creditors", "Net", "Vendor"):
            data_ws.cell(row=r_, column=2).font = Font(bold=True)

    data_ws.sheet_view.showGridLines = False
    data_ws.freeze_panes = "C4"
    data_ws.column_dimensions["A"].width = 3
    for col in "BCDEFGHIJKLMNOP":
        data_ws.column_dimensions[col].width = 15
    note = data_ws.cell(row=max_row_copy + 2, column=2,
                         value=f"Snapshot copied from CFS 08-07-26.xlsx ({sheet_name}).")
    note.font = Font(italic=True, size=9, color=GREY_MUTED)
    data_ws.page_setup.orientation = "landscape"
    data_ws.page_setup.fitToWidth = 1
    data_ws.page_setup.fitToHeight = 0
    data_ws.sheet_properties.pageSetUpPr.fitToPage = True
    return data_ws

copy_source_sheet(src_cur, "Source Data 8-Jul", 41)
copy_source_sheet(src_pri, "Source Data 12-Jun", 35)

CUR = "'Source Data 8-Jul'!"
PRI = "'Source Data 12-Jun'!"

# Row map — current sheet ("Source Data 8-Jul")
CUR_INV, CUR_CUST, CUR_CRED, CUR_NET = 4, 5, 6, 7
# Customer pivot (Outstanding<180=C, Outstanding>180=D, OutstandingTotal=E)
CUR_PIVOT_CUST = {"Machines": 14, "Ink": 13, "Head": 15, "Spare Parts": 17, "Others": 16}
# Vendor pivot (same column meaning)
CUR_PIVOT_VEND = {"Machines": 31, "Ink": 30, "Head": 34, "Spare Parts": 33}

# Row map — prior sheet ("Source Data 12-Jun")
PRI_INV, PRI_CUST, PRI_CRED, PRI_NET = 3, 4, 5, 6

SEG_COL = {"Machines": "C", "Ink": "D", "Head": "E", "Spare Parts": "F", "Others": "G"}

# ================= 2. Summary sheet =================
ws = wb.create_sheet("Summary", 0)
ws.sheet_view.showGridLines = False

for col, w in {"A": 2, "B": 32, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 13,
               "I": 13, "J": 13, "K": 13, "L": 13, "M": 13, "N": 13}.items():
    ws.column_dimensions[col].width = w

def merge_text(cell_range, text_or_formula, size=11, bold=False, color=NAVY_TITLE,
               align="left", italic=False):
    ws.merge_cells(cell_range)
    top_left = cell_range.split(":")[0]
    c = ws[top_left]
    c.value = text_or_formula
    c.font = Font(size=size, bold=bold, color=color, italic=italic, name="Calibri")
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    return c

thin = Side(style="thin", color=GREY_GRID)
box = Border(top=thin, bottom=thin, left=thin, right=thin)

def table_header(row, col_start, headers):
    for i, htext in enumerate(headers):
        c = ws.cell(row=row, column=col_start + i, value=htext)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=NAVY_TITLE)
        c.alignment = Alignment(horizontal="center")
        c.border = box

# ---------------- Title block ----------------
merge_text("B2:N2", "Working Capital Review", size=22, bold=True, color=NAVY_TITLE, align="center")
merge_text("B3:N3", "12-Jun-2026  vs  8-Jul-2026", size=12, color=GREY_MUTED,
           align="center", italic=True)

headline_formula = (
    f'="Net Working Capital up Rs. "&TEXT(({CUR}H{CUR_NET}-{PRI}H{PRI_NET})/10000000,"0.0")'
    f'&" Crore (+"&TEXT(({CUR}H{CUR_NET}-{PRI}H{PRI_NET})/{PRI}H{PRI_NET},"0%")&") in 26 days"'
)
merge_text("B5:N5", headline_formula, size=15, bold=True, color=RED, align="center")

# ================= At a Glance table =================
r = 7
merge_text(f"B{r}:E{r}", "At a Glance (Rs. Crore)", size=13, bold=True)
r += 1
table_header(r, 2, ["Item", "12-Jun-26", "8-Jul-26", "Change"])
glance_header_row = r
r += 1
glance_first_row = r
rows_spec = [
    ("Inventory", CUR_INV, PRI_INV),
    ("Debtors", CUR_CUST, PRI_CUST),
    ("Creditors", CUR_CRED, PRI_CRED),
    ("Net Working Capital", CUR_NET, PRI_NET),
]
for label, cur_row, pri_row in rows_spec:
    ws.cell(row=r, column=2, value=label).border = box
    ws.cell(row=r, column=3, value=f"={PRI}H{pri_row}/10000000").border = box
    ws.cell(row=r, column=4, value=f"={CUR}H{cur_row}/10000000").border = box
    ws.cell(row=r, column=5, value=f"=D{r}-C{r}").border = box
    for col in (3, 4, 5):
        cell = ws.cell(row=r, column=col)
        cell.number_format = '0.0" Cr"'
        cell.alignment = Alignment(horizontal="right")
    if label == "Net Working Capital":
        for col in range(2, 6):
            ws.cell(row=r, column=col).font = Font(bold=True)
    r += 1
glance_last_row = r - 1

# ================= Segment-wise Detail table (self-contained) =================
r = glance_last_row + 2
merge_text(f"B{r}:K{r}", "Segment-wise Detail (Rs. Crore) - 8-Jul-2026 vs 12-Jun-2026", size=13, bold=True)
r += 1
table_header(r, 2, ["Segment", "Inventory (8-Jul)", "Debtors (8-Jul)", "Creditors (8-Jul)", "Net WC (8-Jul)",
                     "Inventory (12-Jun)", "Debtors (12-Jun)", "Creditors (12-Jun)", "Net WC (12-Jun)",
                     "Change in Net WC"])
seg_header_row = r
r += 1
seg_first_row = r
seg_row_of = {}
for label, col in SEG_COL.items():
    seg_row_of[label] = r
    ws.cell(row=r, column=2, value=label).border = box
    ws.cell(row=r, column=3, value=f"={CUR}{col}{CUR_INV}/10000000").border = box
    ws.cell(row=r, column=4, value=f"={CUR}{col}{CUR_CUST}/10000000").border = box
    ws.cell(row=r, column=5, value=f"={CUR}{col}{CUR_CRED}/10000000").border = box
    ws.cell(row=r, column=6, value=f"={CUR}{col}{CUR_NET}/10000000").border = box
    ws.cell(row=r, column=7, value=f"={PRI}{col}{PRI_INV}/10000000").border = box
    ws.cell(row=r, column=8, value=f"={PRI}{col}{PRI_CUST}/10000000").border = box
    ws.cell(row=r, column=9, value=f"={PRI}{col}{PRI_CRED}/10000000").border = box
    ws.cell(row=r, column=10, value=f"={PRI}{col}{PRI_NET}/10000000").border = box
    ws.cell(row=r, column=11,
            value=f"=({CUR}{col}{CUR_NET}-{PRI}{col}{PRI_NET})/10000000").border = box
    for c in range(3, 12):
        cell = ws.cell(row=r, column=c)
        cell.number_format = '+0.0" Cr";-0.0" Cr"' if c == 11 else '0.0" Cr"'
        cell.alignment = Alignment(horizontal="right")
    r += 1
seg_last_row = r - 1
# Total row
ws.cell(row=r, column=2, value="Total").font = Font(bold=True)
ws.cell(row=r, column=2).border = box
ws.cell(row=r, column=3, value=f"={CUR}H{CUR_INV}/10000000").border = box
ws.cell(row=r, column=4, value=f"={CUR}H{CUR_CUST}/10000000").border = box
ws.cell(row=r, column=5, value=f"={CUR}H{CUR_CRED}/10000000").border = box
ws.cell(row=r, column=6, value=f"={CUR}H{CUR_NET}/10000000").border = box
ws.cell(row=r, column=7, value=f"={PRI}H{PRI_INV}/10000000").border = box
ws.cell(row=r, column=8, value=f"={PRI}H{PRI_CUST}/10000000").border = box
ws.cell(row=r, column=9, value=f"={PRI}H{PRI_CRED}/10000000").border = box
ws.cell(row=r, column=10, value=f"={PRI}H{PRI_NET}/10000000").border = box
ws.cell(row=r, column=11, value=f"=({CUR}H{CUR_NET}-{PRI}H{PRI_NET})/10000000").border = box
for c in range(3, 12):
    cell = ws.cell(row=r, column=c)
    cell.number_format = '+0.0" Cr";-0.0" Cr"' if c == 11 else '0.0" Cr"'
    cell.alignment = Alignment(horizontal="right")
    cell.font = Font(bold=True)
seg_total_row = r
r += 1

left_col_end = seg_total_row
r = left_col_end + 3

# ---------------- Page break before the narrative sections ----------------
takeaways_row = r
ws.row_breaks.append(Break(id=r - 1))

def wide_table_header(row, headers):
    for col_start, col_end, htext in headers:
        if col_end and col_end != col_start:
            ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
        c = ws.cell(row=row, column=col_start, value=htext)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=NAVY_TITLE)
        c.alignment = Alignment(horizontal="center")
        for cc in range(col_start, (col_end or col_start) + 1):
            ws.cell(row=row, column=cc).border = box

def data_row_wide(row, cells):
    for col_start, col_end, value, num_fmt, bold in cells:
        if col_end and col_end != col_start:
            ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
        c = ws.cell(row=row, column=col_start, value=value)
        c.border = box
        if num_fmt:
            c.number_format = num_fmt
            c.alignment = Alignment(horizontal="right")
        else:
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        if bold:
            c.font = Font(bold=True)
        for cc in range(col_start, (col_end or col_start) + 1):
            ws.cell(row=row, column=cc).border = box

# ================= Table A: Key Takeaways - Net Working Capital Movement =================
merge_text(f"B{r}:N{r}", "Key Takeaways: Net Working Capital Movement (Rs. Crore)", size=13, bold=True)
r += 1
wide_table_header(r, [(2, None, "Segment"), (3, None, "Now (8-Jul-26)"), (4, None, "Before (12-Jun-26)"),
                       (5, None, "Change"), (6, 14, "Takeaway")])
r += 1

tableA_rows = [
    ("Overall", "H", "Cash conversion is still deteriorating - a further increase in just 26 days.", True),
    ("Ink", "D", "The single biggest driver of the overall increase, by a wide margin.", False),
    ("Head", "E", "Slightly decreased this period - a small offsetting factor, not a driver.", False),
    ("Machines", "C", "A slight increase this period - and receivables risk here is now the highest of the five segments (see table below).", False),
    ("Spare Parts", "F", "Broadly stable - small improvement, still net positive.", False),
    ("Others", "G", "A smaller driver this period - still mostly credit notes/advances.", False),
]
for label, col, takeaway, bold in tableA_rows:
    data_row_wide(r, [
        (2, None, label, None, bold),
        (3, None, f"={CUR}{col}{CUR_NET}/10000000", '0.0" Cr"', bold),
        (4, None, f"={PRI}{col}{PRI_NET}/10000000", '0.0" Cr"', bold),
        (5, None, f"=C{r}-D{r}", '+0.0" Cr";-0.0" Cr"', bold),
        (6, 14, takeaway, None, bold),
    ])
    r += 1
r += 1

# ================= Table B: Key Takeaways - Receivables Ageing Risk =================
merge_text(f"B{r}:N{r}", "Key Takeaways: Receivables Ageing Risk", size=13, bold=True)
r += 1
wide_table_header(r, [(2, None, "Segment"), (3, None, "Aged Receivables >180d (Cr)"),
                       (4, None, "Total Receivables (Cr)"), (5, None, "% Overdue"),
                       (6, 14, "Risk")])
r += 1

tableB_rows = [
    ("Machines", "Highest bad-debt risk - 63% of receivables now aged beyond 180 days."),
    ("Head", "Second-highest risk - 28% aged beyond 180 days."),
    ("Others", "Elevated risk - 21% aged beyond 180 days."),
    ("Spare Parts", "Below-average risk - 10% aged beyond 180 days."),
    ("Ink", "Lowest risk despite the largest absolute exposure after Machines - only 6% aged beyond 180 days."),
]
for label, risk_text in tableB_rows:
    pivot_row = CUR_PIVOT_CUST[label]
    data_row_wide(r, [
        (2, None, label, None, False),
        (3, None, f"={CUR}D{pivot_row}/10000000", '0.0" Cr"', False),
        (4, None, f"={CUR}E{pivot_row}/10000000", '0.0" Cr"', False),
        (5, None, f"={CUR}D{pivot_row}/{CUR}E{pivot_row}", "0%", False),
        (6, 14, risk_text, None, False),
    ])
    r += 1
r += 1

# ================= Table C: Action Plan =================
merge_text(f"B{r}:N{r}", "Action Plan - Target: Q2 FY27 (Jul-Sep 2026)", size=13, bold=True)
r += 1
wide_table_header(r, [(2, None, "Segment"), (3, 4, "Action"), (5, None, "Now"),
                       (6, None, "Target (30-Sep-26)"), (7, 14, "Basis")])
r += 1

# (segment, action, now_formula, now_fmt, target_kind, basis, bold)
# target_kind: "reduce_17_5" | "recover_0.30" | "recover_0.25" | literal string | None (dash)
tableC_rows = [
    ("Ink", "Reduce inventory", f"={CUR}D{CUR_INV}/10000000", '0.0" Cr"',
     "reduce_17_5", "17.5% cut (midpoint of 15-20%)", False),
    ("Ink", "Recover aged receivables", f"={CUR}D{CUR_PIVOT_CUST['Ink']}/10000000", '0.0" Cr"',
     "recover_0.30", "30% of aged >180d", False),
    ("Head", "Reduce inventory", f"={CUR}E{CUR_INV}/10000000", '0.0" Cr"',
     "reduce_17_5", "17.5% cut (midpoint of 15-20%)", False),
    ("Head", "Recover aged receivables", f"={CUR}D{CUR_PIVOT_CUST['Head']}/10000000", '0.0" Cr"',
     "recover_0.30", "30% of aged >180d", False),
    ("Machines", "Freeze dispatch + recover aged dues", f"={CUR}D{CUR_PIVOT_CUST['Machines']}/10000000", '0.0" Cr"',
     "recover_0.25", "25% of aged >180d", False),
    ("Spare Parts", "Maintain discipline",
     f"={CUR}D{CUR_PIVOT_CUST['Spare Parts']}/{CUR}E{CUR_PIVOT_CUST['Spare Parts']}", "0%",
     "-", "Below-average risk - maintain discipline", False),
    ("Others", "Investigate the swing", f"=({CUR}G{CUR_NET}-{PRI}G{PRI_NET})/10000000", '0.0" Cr swing"',
     "-", "Confirm not a reconciliation issue", False),
    ("Overall", "Reduce Net Working Capital", f"={CUR}H{CUR_NET}/10000000", '0.0" Cr"',
     "45 - 50", "Policy target", True),
]
for segment, action, now_formula, now_fmt, target_kind, basis, bold in tableC_rows:
    ws.cell(row=r, column=2, value=segment)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    ws.cell(row=r, column=3, value=action)
    ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="center")
    c_now = ws.cell(row=r, column=5, value=now_formula)
    c_now.number_format = now_fmt
    c_now.alignment = Alignment(horizontal="right")
    if target_kind == "reduce_17_5":
        c_target = ws.cell(row=r, column=6, value=f"=E{r}*0.825")
        c_target.number_format = '0.0" Cr"'
        c_target.alignment = Alignment(horizontal="right")
    elif target_kind == "recover_0.30":
        c_target = ws.cell(row=r, column=6, value=f'="Recover "&TEXT(E{r}*0.3,"0.0")&" Cr"')
        c_target.alignment = Alignment(horizontal="right")
    elif target_kind == "recover_0.25":
        c_target = ws.cell(row=r, column=6, value=f'="Recover "&TEXT(E{r}*0.25,"0.0")&" Cr"')
        c_target.alignment = Alignment(horizontal="right")
    else:
        c_target = ws.cell(row=r, column=6, value=target_kind)
        c_target.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=14)
    ws.cell(row=r, column=7, value=basis)
    ws.cell(row=r, column=7).alignment = Alignment(wrap_text=True, vertical="center")
    for cc in range(2, 15):
        ws.cell(row=r, column=cc).border = box
        if bold:
            ws.cell(row=r, column=cc).font = Font(bold=True)
    r += 1

last_row = r + 2

# ================= Print setup =================
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.2, footer=0.2)
ws.print_area = f"A1:N{last_row}"
ws.print_options.horizontalCentered = True

wb.save(OUT)
print("Saved:", OUT)
print("left_col_end", left_col_end, "takeaways_row", takeaways_row, "last_row", last_row)
